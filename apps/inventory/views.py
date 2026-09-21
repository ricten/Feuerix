from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.core.crud import abschnitt, knopf

from .models import Gegenstand, Inventur, Inventurposition, Verleih
from .pdf import etiketten_pdf, leihschein_pdf


def _pruefen(request, modul, aktion):
    if request.verein is None or not request.rechte.darf(modul, aktion):
        raise PermissionDenied


# ---------------------------------------------------------------- Gegenstand
def gegenstand_kontext(request, g):
    aktionen = []
    if request.rechte.darf("inventar", "view"):
        aktionen.append(knopf("Etikett drucken", reverse("gegenstand_etikett", args=[g.pk])))
    if request.rechte.darf("verleih", "add") and g.verleihbar:
        aktionen.append(knopf("Verleih / Reservierung anlegen", reverse("verleih_add") + f"?gegenstand={g.pk}",
                              stil="primary"))
    hinweise = []
    akt = g.aktuell_verliehen
    if akt:
        hinweise.append(f"Aktuell verliehen an {akt.wer} bis {akt.bis:%d.%m.%Y}.")
    if g.garantie_bis and g.garantie_bis < date.today():
        hinweise.append("Garantie abgelaufen.")
    abschnitte = []
    if request.rechte.darf("verleih", "view"):
        abschnitte.append(abschnitt(request, "Verleihhistorie", g.verleihe.all(),
                                    ("von", "bis", ("wer", "Entleiher"), "status", "veranstaltung")))
    return {"aktionen": aktionen, "hinweise": hinweise, "abschnitte": abschnitte}


# ---------------------------------------------------------------- Verleih
def verleih_kontext(request, v):
    rt, aktionen = request.rechte, [knopf("Leihschein (PDF)", reverse("verleih_leihschein", args=[v.pk]))]
    if rt.darf("verleih", "change"):
        if v.status == "reserviert":
            aktionen.append(knopf("Ausgeben", reverse("verleih_ausgeben", args=[v.pk]), post=True, stil="success"))
            aktionen.append(knopf("Reservierung stornieren", reverse("verleih_stornieren", args=[v.pk]), post=True,
                                  stil="outline-danger", bestaetigung="Reservierung stornieren?"))
        if v.status == "ausgegeben":
            aktionen.append(knopf("Rückgabe – in Ordnung", reverse("verleih_rueckgabe", args=[v.pk]), post=True,
                                  stil="success", felder={"zustand": v.gegenstand.zustand}))
            aktionen.append(knopf("Rückgabe – defekt", reverse("verleih_rueckgabe", args=[v.pk]), post=True,
                                  stil="outline-danger", felder={"zustand": "defekt"},
                                  bestaetigung="Gegenstand als defekt markieren?"))
    hinweise = []
    if v.ueberfaellig:
        hinweise.append("Rückgabe überfällig!")
    if v.status == "zurueckgegeben" and v.kaution and not v.kaution_zurueckgezahlt:
        hinweise.append(f"Kaution ({v.kaution} €) noch nicht als zurückgezahlt markiert.")
    return {"aktionen": aktionen, "hinweise": hinweise}


def _verleih(request, pk):
    _pruefen(request, "verleih", "change")
    return get_object_or_404(Verleih, pk=pk, verein=request.verein)


@login_required
@require_POST
def verleih_ausgeben(request, pk):
    v = _verleih(request, pk)
    if v.status != "reserviert":
        messages.error(request, "Nur Reservierungen können ausgegeben werden.")
    elif v.gegenstand.zustand in ("defekt", "ausgesondert"):
        messages.error(request, "Gegenstand ist defekt bzw. ausgesondert.")
    else:
        v.status, v.ausgegeben_am = "ausgegeben", timezone.now()
        v.ausgegeben_von = request.user.get_username()
        v.zustand_bei_ausgabe = v.gegenstand.zustand
        v.save()
        messages.success(request, "Ausgegeben.")
    return redirect("verleih_detail", pk=v.pk)


@login_required
@require_POST
def verleih_rueckgabe(request, pk):
    v = _verleih(request, pk)
    zustand = request.POST.get("zustand", "")
    if v.status != "ausgegeben":
        messages.error(request, "Nur ausgegebene Gegenstände können zurückgenommen werden.")
    elif zustand not in dict(Gegenstand.ZUSTAND):
        messages.error(request, "Ungültiger Zustand.")
    else:
        v.status, v.zurueckgegeben_am, v.zustand_bei_rueckgabe = "zurueckgegeben", timezone.now(), zustand
        v.save()
        if v.gegenstand.zustand != zustand:
            v.gegenstand.zustand = zustand
            v.gegenstand.save()
        messages.success(request, "Rückgabe gebucht." + (" Bitte Kaution zurückzahlen." if v.kaution else ""))
    return redirect("verleih_detail", pk=v.pk)


@login_required
@require_POST
def verleih_stornieren(request, pk):
    v = _verleih(request, pk)
    if v.status == "reserviert":
        v.status = "storniert"
        v.save()
    return redirect("verleih_detail", pk=v.pk)


@login_required
def verleih_leihschein(request, pk):
    _pruefen(request, "verleih", "view")
    v = get_object_or_404(Verleih, pk=pk, verein=request.verein)
    r = HttpResponse(leihschein_pdf(v), content_type="application/pdf")
    r["Content-Disposition"] = f'inline; filename="leihschein-{v.pk}.pdf"'
    return r


# ---------------------------------------------------------------- Inventur
def inventur_kontext(request, inv):
    z = inv.zaehlung()
    aktionen = []
    if request.rechte.darf("inventur", "change") and inv.status == "laufend":
        aktionen.append(knopf("Inventur abschließen", reverse("inventur_abschliessen", args=[inv.pk]), post=True,
                              stil="success", bestaetigung="Inventur abschließen? Danach sind die Ergebnisse fixiert."))
    gesamt = sum(z.values())
    hinweise = [f"{gesamt} Gegenstände · ✓ {z['gefunden']} gefunden · ⚠ {z['nicht_gefunden']} nicht gefunden · "
                f"⚠ {z['beschaedigt']} beschädigt · offen {z['offen']}"]
    zeilen = []
    kann = request.rechte.darf("inventur", "change") and inv.status == "laufend"
    filt = request.GET.get("zeige", "")
    for p in inv.positionen.all():
        if filt and p.ergebnis != filt:
            continue
        akt = []
        if kann:
            for erg, label, stil in (("gefunden", "✓", "success"), ("nicht_gefunden", "✗", "danger"),
                                     ("beschaedigt", "⚠", "warning")):
                akt.append(knopf(label, reverse("inventurposition_setzen", args=[p.pk]), post=True,
                                 stil=("" if p.ergebnis != erg else "") + stil, felder={"ergebnis": erg}))
        zeilen.append({"url": None, "zellen": [p.inventarnummer, p.bezeichnung, p.standort_text,
                                                p.get_ergebnis_display()], "aktionen": akt})
    return {"aktionen": aktionen, "hinweise": hinweise, "abschnitte": [
        {"titel": "Positionen" + (f" (Filter: {filt})" if filt else ""),
         "spalten": ["Nr.", "Bezeichnung", "Standort (Soll)", "Ergebnis"], "zeilen": zeilen, "add_url": None,
         "mit_aktionen": True}]}


@login_required
@require_POST
def inventurposition_setzen(request, pk):
    _pruefen(request, "inventur", "change")
    p = get_object_or_404(Inventurposition, pk=pk, verein=request.verein)
    if p.inventur.status != "laufend":
        raise PermissionDenied
    erg = request.POST.get("ergebnis")
    if erg in dict(Inventurposition.ERGEBNIS):
        p.ergebnis = erg
        p.save(update_fields=["ergebnis", "geaendert"])
    return redirect("inventur_detail", pk=p.inventur_id)


@login_required
@require_POST
def inventur_abschliessen(request, pk):
    _pruefen(request, "inventur", "change")
    inv = get_object_or_404(Inventur, pk=pk, verein=request.verein)
    if inv.status == "laufend":
        inv.status, inv.abgeschlossen_am = "abgeschlossen", date.today()
        inv.save()
        messages.success(request, "Inventur abgeschlossen.")
    return redirect("inventur_detail", pk=inv.pk)


# ---------------------------------------------------------------- Import / Etiketten / Scan
@login_required
def gegenstand_import(request):
    from django.shortcuts import render

    from . import importer
    if request.verein is None or not request.rechte.darf("inventar", "add"):
        raise PermissionDenied
    bericht = None
    if request.method == "POST":
        datei = request.FILES.get("datei")
        if not datei:
            messages.error(request, "Bitte eine Datei auswählen.")
        else:
            try:
                bericht = importer.importieren(
                    request.verein, datei.name, datei.read(), testlauf=bool(request.POST.get("testlauf")),
                    aktualisieren=bool(request.POST.get("aktualisieren")), neu_anlegen=bool(request.POST.get("neu_anlegen")))
            except ValueError as e:
                messages.error(request, str(e))
    return render(request, "inventory/import.html", {"titel": "Inventar importieren", "bericht": bericht,
                                                      "post": request.POST if request.method == "POST" else {}})


@login_required
def gegenstand_import_vorlage(request):
    if request.verein is None or not request.rechte.darf("inventar", "add"):
        raise PermissionDenied
    from io import BytesIO

    from openpyxl import Workbook

    from .tabellen import SPALTEN_ANZEIGE
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventar"
    ws.append([label for _, label in SPALTEN_ANZEIGE])
    beispiel = {"bezeichnung": "Beamer Epson EB-X05", "kategorie": "Technik", "standort": "Lager",
               "hersteller": "Epson", "modell": "EB-X05", "anschaffungsdatum": "01.03.2022",
               "anschaffungspreis": "450,00", "zustand": "gut", "verleihbar": "ja", "kaution": "50,00"}
    ws.append([beispiel.get(f, "") for f, _ in SPALTEN_ANZEIGE])
    for i, _ in enumerate(SPALTEN_ANZEIGE, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 18
    hilfe = wb.create_sheet("Hinweise")
    for z in ["Pflichtspalte: Bezeichnung.",
              "Leere Inventarnummer = automatisch vergeben (INV-000001 ...). Vorhandene Nummer = bestehender "
              "Gegenstand wird aktualisiert.",
              "Leere Zellen überschreiben keine vorhandenen Daten.",
              "Datum: TT.MM.JJJJ. Zustand: neu, gut, gebrauchsspuren, defekt, ausgesondert. Verleihbar: ja/nein.",
              "Kategorie/Standort müssen existieren (oder Option „unbekannte anlegen“ beim Import).",
              "Zeile 2 ist ein Beispiel und sollte gelöscht werden."]:
        hilfe.append([z])
    buf = BytesIO()
    wb.save(buf)
    r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = 'attachment; filename="inventar-import-vorlage.xlsx"'
    return r


def _scan_url(request):
    def bauen(inventarnummer):
        return request.build_absolute_uri(reverse("gegenstand_scan", args=[inventarnummer]))
    return bauen


@login_required
def gegenstand_etikett(request, pk):
    _pruefen(request, "inventar", "view")
    g = get_object_or_404(Gegenstand, pk=pk, verein=request.verein)
    try:
        anzahl = max(1, min(24, int(request.GET.get("anzahl", 1))))
    except ValueError:
        anzahl = 1
    pdf = etiketten_pdf([g] * anzahl, _scan_url(request))
    r = HttpResponse(pdf, content_type="application/pdf")
    r["Content-Disposition"] = f'inline; filename="etikett-{g.inventarnummer}.pdf"'
    return r


@login_required
def gegenstand_etiketten(request):
    _pruefen(request, "inventar", "view")
    qs = Gegenstand.objects.filter(verein=request.verein).exclude(zustand="ausgesondert").order_by("inventarnummer")
    if not qs.exists():
        messages.info(request, "Kein Inventar vorhanden.")
        return redirect("gegenstand_list")
    pdf = etiketten_pdf(qs, _scan_url(request))
    r = HttpResponse(pdf, content_type="application/pdf")
    r["Content-Disposition"] = 'inline; filename="inventar-etiketten.pdf"'
    return r


@login_required
def gegenstand_scan(request, inventarnummer):
    """Ziel des QR-Codes auf dem Etikett: findet den Gegenstand über die Inventarnummer und führt direkt in den
    (vorausgefüllten) Verleih-Start - z. B. wenn beim Ausleihen per Handy gescannt wird."""
    if request.verein is None:
        return redirect("verein_waehlen")
    g = get_object_or_404(Gegenstand, verein=request.verein, inventarnummer=inventarnummer)
    if not request.rechte.darf("inventar", "view"):
        raise PermissionDenied
    if not g.verleihbar or g.zustand in ("defekt", "ausgesondert"):
        messages.warning(request, "Dieser Gegenstand ist nicht als verleihbar markiert bzw. defekt/ausgesondert.")
        return redirect("gegenstand_detail", pk=g.pk)
    if not request.rechte.darf("verleih", "add"):
        return redirect("gegenstand_detail", pk=g.pk)
    return redirect(reverse("verleih_add") + f"?gegenstand={g.pk}")
