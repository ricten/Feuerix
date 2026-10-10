from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.core.models import Rolle, Verein, Zugang
from apps.events.models import Veranstaltung, Wahlergebnis
from apps.members.models import Mitglied

from . import services
from .client import OpenSlidesFehler
from .models import OpenSlidesVerbindung


class VerbindungOderNoneTests(TestCase):
    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")

    def test_keine_verbindung_angelegt(self):
        self.assertIsNone(services.verbindung_oder_none(self.v))

    def test_unvollstaendige_verbindung(self):
        OpenSlidesVerbindung.objects.create(verein=self.v, url="https://os.example.org")  # ohne Benutzer/Passwort
        self.assertIsNone(services.verbindung_oder_none(self.v))

    def test_vollstaendige_verbindung(self):
        v = OpenSlidesVerbindung.objects.create(verein=self.v, url="https://os.example.org",
                                                benutzername="sync", passwort="geheim")
        self.assertEqual(services.verbindung_oder_none(self.v), v)


class MitgliedAnonymisierenTests(TestCase):
    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.m = Mitglied.objects.create(verein=self.v, vorname="Erika", nachname="Muster", mitgliedsnummer=7,
                                         openslides_user_id=42, openslides_username="erika.muster")

    def test_ohne_openslides_konto_passiert_nichts(self):
        m2 = Mitglied.objects.create(verein=self.v, vorname="Ohne", nachname="Konto")
        with patch("apps.openslides.services.OSClient") as MockClient:
            services.mitglied_anonymisieren(self.verbindung, m2)
        MockClient.assert_not_called()

    def test_ruft_user_update_mit_anonymisierten_daten_auf(self):
        with patch("apps.openslides.services.OSClient") as MockClient:
            instanz = MockClient.return_value
            services.mitglied_anonymisieren(self.verbindung, self.m)
        instanz.login.assert_called_once()
        name, daten_liste = instanz.action.call_args.args
        self.assertEqual(name, "user.update")
        daten = daten_liste[0]
        self.assertEqual(daten["id"], 42)
        self.assertEqual(daten["first_name"], "Anonymisiert")
        self.assertEqual(daten["last_name"], "#7")
        self.assertEqual(daten["email"], "")
        self.assertEqual(daten["is_active"], False)
        self.assertNotIn("muster", daten["username"].lower())

    def test_fehler_wird_weitergereicht(self):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.login.side_effect = OpenSlidesFehler("kaputt")
            with self.assertRaises(OpenSlidesFehler):
                services.mitglied_anonymisieren(self.verbindung, self.m)

    def test_ohne_mitgliedsnummer_faellt_auf_pk_zurueck(self):
        self.m.mitgliedsnummer = None
        self.m.save()
        with patch("apps.openslides.services.OSClient") as MockClient:
            instanz = MockClient.return_value
            services.mitglied_anonymisieren(self.verbindung, self.m)
        daten = instanz.action.call_args.args[1][0]
        self.assertEqual(daten["last_name"], f"#{self.m.pk}")


def _wahl_antwort():
    """Simulierte Autoupdate-Antwort: eine Wahl ("1. Vorsitzender") mit einem abgeschlossenen Wahlgang und
    zwei Kandidaten (Max Muster, Erika Musterfrau) - nach dem Datenmodell aus openslides-meta (assignment ->
    poll -> option -> poll_candidate_list -> poll_candidate -> user)."""
    return {
        "meeting/5/assignment_ids": [12],
        "assignment/12/title": "1. Vorsitzender",
        "assignment/12/poll_ids": [45],
        "poll/45/title": "Wahlgang 1",
        "poll/45/state": "published",
        "poll/45/option_ids": [100, 101],
        "option/100/yes": "8.000000", "option/100/no": "1.000000", "option/100/abstain": "1.000000",
        "option/100/content_object_id": "poll_candidate_list/7",
        "poll_candidate_list/7/poll_candidate_ids": [200],
        "poll_candidate/200/user_id": 42,
        "user/42/first_name": "Max", "user/42/last_name": "Muster",
        "option/101/yes": "3.000000", "option/101/no": "5.000000", "option/101/abstain": "2.000000",
        "option/101/content_object_id": "poll_candidate_list/8",
        "poll_candidate_list/8/poll_candidate_ids": [201],
        "poll_candidate/201/user_id": 43,
        "user/43/first_name": "Erika", "user/43/last_name": "Musterfrau",
    }


class WahlergebnisseAbrufenTests(TestCase):
    """Rückfluss der Wahlergebnisse aus OpenSlides (assignment/poll/option/poll_candidate/user) ins Protokoll -
    die eigentliche HTTP-Abfrage wird über OSClient.abfragen gemockt, da keine echte Instanz erreichbar ist."""

    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)

    def test_ohne_meeting_id_wird_abgelehnt(self):
        self.ver.openslides_meeting_id = None
        self.ver.save()
        with self.assertRaises(OpenSlidesFehler):
            services.wahlergebnisse_abrufen(self.verbindung, self.ver)

    def test_speichert_ergebnis_je_wahlgang(self):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.return_value = _wahl_antwort()
            n = services.wahlergebnisse_abrufen(self.verbindung, self.ver)
        self.assertEqual(n, 1)
        w = Wahlergebnis.objects.get(verein=self.v, veranstaltung=self.ver)
        self.assertEqual(w.amt, "1. Vorsitzender")
        self.assertEqual(w.wahlgang, "Wahlgang 1")
        self.assertIn("Max Muster: 8 Ja, 1 Nein, 1 Enthaltung", w.ergebnis)
        self.assertIn("Erika Musterfrau: 3 Ja, 5 Nein, 2 Enthaltung", w.ergebnis)
        self.assertEqual(w.art, "wahl")
        self.assertEqual(w.openslides_poll_id, 45)

    def test_nicht_abgeschlossene_wahl_wird_uebersprungen(self):
        antwort = _wahl_antwort()
        antwort["poll/45/state"] = "started"
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.return_value = antwort
            n = services.wahlergebnisse_abrufen(self.verbindung, self.ver)
        self.assertEqual(n, 0)
        self.assertFalse(Wahlergebnis.objects.filter(veranstaltung=self.ver).exists())

    def test_erneuter_abruf_ersetzt_alte_ergebnisse(self):
        Wahlergebnis.objects.create(verein=self.v, veranstaltung=self.ver, amt="Veraltet", wahlgang="",
                                    ergebnis="nicht mehr aktuell")
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.return_value = _wahl_antwort()
            services.wahlergebnisse_abrufen(self.verbindung, self.ver)
        ergebnisse = Wahlergebnis.objects.filter(veranstaltung=self.ver)
        self.assertEqual(ergebnisse.count(), 1)
        self.assertFalse(ergebnisse.filter(amt="Veraltet").exists())


class VeranstaltungWahlergebnisseViewTests(TestCase):
    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        OpenSlidesVerbindung.objects.create(verein=self.v, url="https://os.example.org", benutzername="sync",
                                            passwort="geheim", aktiv=True)
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", password="pw-Test-12345")
        Zugang.objects.create(verein=self.v, user=self.admin, rolle=Rolle.objects.get(verein=self.v,
                                                                                       name="Superadministrator"))
        self.client.login(username="admin", password="pw-Test-12345")

    def test_knopf_erscheint_wenn_meeting_verknuepft(self):
        r = self.client.get(reverse("veranstaltung_detail", args=[self.ver.pk]))
        self.assertContains(r, "Anwesenheit, Wahlen und Abstimmungen aus OpenSlides übernehmen")

    def test_ohne_meeting_verknuepfung_kein_knopf(self):
        self.ver.openslides_meeting_id = None
        self.ver.save()
        r = self.client.get(reverse("veranstaltung_detail", args=[self.ver.pk]))
        self.assertNotContains(r, "Anwesenheit, Wahlen und Abstimmungen aus OpenSlides übernehmen")

    def test_post_uebernimmt_ergebnisse_und_zeigt_abschnitt(self):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.return_value = _wahl_antwort()
            r = self.client.post(reverse("veranstaltung_wahlergebnisse", args=[self.ver.pk]), follow=True)
        self.assertContains(r, "1 Wahlergebnis(se), 0 Abstimmung(en) und 0 Anwesende aus OpenSlides übernommen.")
        r2 = self.client.get(reverse("veranstaltung_detail", args=[self.ver.pk]))
        self.assertContains(r2, "Wahlergebnisse (aus OpenSlides)")
        self.assertContains(r2, "1. Vorsitzender")


class SsrfSchutzTests(TestCase):
    """Die OpenSlides-Adresse darf nicht auf interne Dienste (SSRF) zeigen - siehe
    apps.core.util.pruefe_oeffentliche_adresse."""

    def test_interne_adresse_wird_bei_voller_pruefung_abgelehnt(self):
        from django.core.exceptions import ValidationError
        v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        verbindung = OpenSlidesVerbindung(verein=v, url="http://169.254.169.254/", benutzername="bot")
        with self.assertRaises(ValidationError):
            verbindung.clean()


def _antrag_antwort():
    """Simulierte Autoupdate-Antwort fuer Antraege: Antrag A1 mit einer abgeschlossenen Abstimmung (Ja/Nein/
    Enthaltung auf EINER Option ohne Text), Status "beschlossen", dazu eine noch laufende Abstimmung."""
    return {
        "meeting/5/motion_ids": [7],
        "motion/7/number": "A1", "motion/7/title": "Beitragserhöhung", "motion/7/state_id": 3,
        "motion_state/3/name": "beschlossen",
        "motion/7/poll_ids": [60, 61],
        "poll/60/title": "Abstimmung", "poll/60/state": "published", "poll/60/option_ids": [300],
        "poll/60/votesvalid": "12.000000", "poll/60/votesinvalid": "0.000000", "poll/60/votescast": "12.000000",
        "option/300/yes": "9.000000", "option/300/no": "2.000000", "option/300/abstain": "1.000000",
        "poll/61/title": "Läuft noch", "poll/61/state": "started", "poll/61/option_ids": [301],
        "option/301/yes": "1.000000", "option/301/no": "0.000000", "option/301/abstain": "0.000000",
    }


def _thema_antwort():
    return {
        "meeting/5/topic_ids": [9],
        "topic/9/title": "Entlastung des Vorstands",
        "topic/9/poll_ids": [70],
        "poll/70/title": "Abstimmung Entlastung", "poll/70/state": "finished", "poll/70/option_ids": [400],
        "option/400/yes": "10", "option/400/no": "0", "option/400/abstain": "2",
    }


class AbstimmungenAbrufenTests(TestCase):
    """Alle Abstimmungen (Antraege, Themen) einer OpenSlides-Versammlung ins Protokoll uebernehmen. Die HTTP-
    Abfragen (Wahlen, Antraege, Themen - in dieser Reihenfolge) werden ueber OSClient.abfragen gemockt."""

    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)

    def _abrufen(self, wahlen=None, antraege=None, themen=None):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = [wahlen or {}, antraege or {}, themen or {}]
            return services.wahlergebnisse_abrufen(self.verbindung, self.ver)

    def test_antrag_wird_mit_stimmen_gesamtzahlen_und_status_gespeichert(self):
        self.assertEqual(self._abrufen(antraege=_antrag_antwort()), 1)
        a = Wahlergebnis.objects.get(veranstaltung=self.ver)
        self.assertEqual(a.art, "abstimmung")
        self.assertEqual(a.amt, "Antrag A1: Beitragserhöhung")
        self.assertEqual(a.wahlgang, "Abstimmung")
        self.assertEqual(a.openslides_poll_id, 60)
        self.assertIn("9 Ja, 2 Nein, 1 Enthaltung", a.ergebnis)
        self.assertIn("Gültige Stimmen: 12", a.ergebnis)
        self.assertIn("Status in OpenSlides: beschlossen", a.ergebnis)
        self.assertNotIn("Option", a.ergebnis)

    def test_laufende_abstimmung_wird_uebersprungen(self):
        self._abrufen(antraege=_antrag_antwort())
        self.assertFalse(Wahlergebnis.objects.filter(openslides_poll_id=61).exists())

    def test_thema_wird_uebernommen(self):
        self.assertEqual(self._abrufen(themen=_thema_antwort()), 1)
        t = Wahlergebnis.objects.get(veranstaltung=self.ver)
        self.assertEqual((t.art, t.amt), ("abstimmung", "Thema: Entlastung des Vorstands"))
        self.assertIn("10 Ja, 0 Nein, 2 Enthaltung", t.ergebnis)

    def test_wahlen_und_abstimmungen_zusammen_und_idempotent(self):
        n = self._abrufen(wahlen=_wahl_antwort(), antraege=_antrag_antwort(), themen=_thema_antwort())
        self.assertEqual(n, 3)
        self.assertEqual(Wahlergebnis.objects.filter(art="wahl").count(), 1)
        self.assertEqual(Wahlergebnis.objects.filter(art="abstimmung").count(), 2)
        self._abrufen(antraege=_antrag_antwort())   # erneuter Abruf ersetzt alles
        self.assertEqual(Wahlergebnis.objects.count(), 1)

    def test_fehlerantwort_in_allen_varianten_wird_gemeldet_statt_verschluckt(self):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = [{}] + [{"message": "kaputt"}] * 4
            with self.assertRaisesMessage(OpenSlidesFehler, "Abstimmungen zu Anträgen konnten nicht abgerufen werden"):
                services.wahlergebnisse_abrufen(self.verbindung, self.ver)

    def test_zahlen_lesbar(self):
        self.assertEqual(services._zahl("8.000000"), "8")
        self.assertEqual(services._zahl("0.500000"), "0,5")
        self.assertEqual(services._zahl("10"), "10")
        self.assertEqual(services._zahl("keine"), "keine")


class ProtokollAbstimmungenViewTests(TestCase):
    def setUp(self):
        from apps.documents.models import Schriftstueck
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        OpenSlidesVerbindung.objects.create(verein=self.v, url="https://os.example.org", benutzername="sync",
                                            passwort="geheim", aktiv=True)
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)
        self.s = Schriftstueck.objects.create(verein=self.v, titel="Protokoll MV", art="protokoll",
                                              veranstaltung=self.ver, text="Beginn der Sitzung.")
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", password="pw-Test-12345")
        Zugang.objects.create(verein=self.v, user=self.admin, rolle=Rolle.objects.get(verein=self.v,
                                                                                       name="Superadministrator"))
        self.client.login(username="admin", password="pw-Test-12345")
        self.url = reverse("schriftstueck_abstimmungen", args=[self.s.pk])

    def _post(self, **antworten):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = [antworten.get("wahlen", {}),
                                                            antworten.get("antraege", {}),
                                                            antworten.get("themen", {}),
                                                            antworten.get("anwesenheit", {})]
            return self.client.post(self.url, follow=True)

    def test_knopf_nur_bei_protokoll_mit_verknuepfter_versammlung(self):
        self.assertContains(self.client.get(reverse("schriftstueck_detail", args=[self.s.pk])),
                            "Anwesenheit, Abstimmungen und Wahlen aus OpenSlides abrufen (als Anlagen)")
        self.ver.openslides_meeting_id = None
        self.ver.save()
        self.assertNotContains(self.client.get(reverse("schriftstueck_detail", args=[self.s.pk])),
                               "aus OpenSlides abrufen (als Anlagen)")

    def test_abruf_meldet_die_angehaengten_anlagen_und_laesst_den_text_unveraendert(self):
        r = self._post(antraege=_antrag_antwort(), wahlen=_wahl_antwort(), anwesenheit=_anwesenheit_antwort())
        self.assertContains(r, "1 Wahlergebnis(se), 1 Abstimmung(en) und 2 Anwesende aus OpenSlides übernommen.")
        self.assertContains(r, "Anlage 1: Teilnehmerliste, Anlage 2: Abstimmungen, Anlage 3: Wahlergebnisse")
        self.s.refresh_from_db()
        self.assertEqual(self.s.text, "Beginn der Sitzung.")

    def test_pdf_und_word_des_protokolls_enthalten_die_anlagen(self):
        self._post(antraege=_antrag_antwort(), wahlen=_wahl_antwort(), anwesenheit=_anwesenheit_antwort())
        import io
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(self.client.get(reverse("schriftstueck_pdf", args=[self.s.pk])).content))
        text = "\n".join(seite.extract_text() for seite in pdf.pages)
        self.assertEqual(len(pdf.pages), 4)   # Protokoll + 3 Anlagen, je eigene Seite
        for erwartet in ("Anlage 1: Teilnehmerliste", "Gast Besucher", "Anlage 2: Abstimmungen", "Antrag A1",
                         "Anlage 3: Wahlergebnisse", "1. Vorsitzender", "Max Muster"):
            self.assertIn(erwartet, text)
        from docx import Document
        word = Document(io.BytesIO(self.client.get(reverse("schriftstueck_docx", args=[self.s.pk])).content))
        absaetze = [p.text for p in word.paragraphs]
        for erwartet in ("Anlage 1: Teilnehmerliste", "Anlage 2: Abstimmungen", "Anlage 3: Wahlergebnisse"):
            self.assertIn(erwartet, absaetze)
        zellen = [z.text for t in word.tables for r in t.rows for z in r.cells]
        self.assertIn("Gast Besucher", zellen)
        self.assertIn("Max Muster", zellen)

    def test_ohne_ergebnisse_warnung_und_text_unveraendert(self):
        r = self._post()
        self.assertContains(r, "Keine abgeschlossenen Wahlen oder Abstimmungen und keine Anwesenden")
        self.s.refresh_from_db()
        self.assertEqual(self.s.text, "Beginn der Sitzung.")

    def test_abgelehnt_ohne_verknuepfung_oder_fremdes_protokoll(self):
        self.ver.openslides_meeting_id = None
        self.ver.save()
        r = self.client.post(self.url, follow=True)
        self.assertContains(r, "Nur für Protokolle")

    def test_nur_mit_rechten(self):
        User = get_user_model()
        User.objects.create_user("anna", password="pw-Test-12345")
        self.client.login(username="anna", password="pw-Test-12345")
        self.assertIn(self.client.post(self.url).status_code, (302, 403))
        self.s.refresh_from_db()
        self.assertEqual(self.s.text, "Beginn der Sitzung.")


def _anwesenheit_antwort():
    """Simulierte Autoupdate-Antwort: zwei als anwesend markierte Personen - ein Mitglied (OpenSlides-Konto 42) und
    ein Gast ohne Mitgliedskonto."""
    return {
        "meeting/5/present_user_ids": [42, 99],
        "user/42/first_name": "Erika", "user/42/last_name": "OS-Name-wird-nicht-gebraucht",
        "user/99/first_name": "Gast", "user/99/last_name": "Besucher",
    }


class AnwesenheitAbrufenTests(TestCase):
    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)
        self.m = Mitglied.objects.create(verein=self.v, vorname="Erika", nachname="Muster", mitgliedsnummer=7,
                                         openslides_user_id=42, openslides_username="erika.muster")

    def _abrufen(self, antwort):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.return_value = antwort
            return services.anwesenheit_abrufen(self.verbindung, self.ver)

    def test_mitglied_wird_zugeordnet_gast_behaelt_namen(self):
        from apps.events.models import Anwesenheit
        self.assertEqual(self._abrufen(_anwesenheit_antwort()), 2)
        mitglied = Anwesenheit.objects.get(openslides_user_id=42)
        self.assertEqual(mitglied.mitglied, self.m)
        self.assertEqual(mitglied.name, "")   # Name wird nicht kopiert (DSGVO-Anonymisierung wirkt)
        self.assertEqual(mitglied.wer, "Erika Muster")
        gast = Anwesenheit.objects.get(openslides_user_id=99)
        self.assertIsNone(gast.mitglied)
        self.assertEqual(gast.wer, "Gast Besucher")

    def test_mitglied_eines_anderen_vereins_wird_nicht_zugeordnet(self):
        from apps.events.models import Anwesenheit
        anderer = Verein.objects.create(name="Anderer e.V.", kuerzel="anderer")
        Mitglied.objects.create(verein=anderer, vorname="Fremd", nachname="Person", mitgliedsnummer=1,
                                openslides_user_id=99)
        self._abrufen(_anwesenheit_antwort())
        self.assertIsNone(Anwesenheit.objects.get(openslides_user_id=99).mitglied)

    def test_erneuter_abruf_ersetzt_den_stand(self):
        from apps.events.models import Anwesenheit
        self._abrufen(_anwesenheit_antwort())
        self._abrufen({"meeting/5/present_user_ids": [42], "user/42/first_name": "Erika"})
        self.assertEqual(Anwesenheit.objects.filter(veranstaltung=self.ver).count(), 1)

    def test_niemand_anwesend_ist_kein_fehler(self):
        from apps.events.models import Anwesenheit
        self._abrufen(_anwesenheit_antwort())
        self.assertEqual(self._abrufen({"meeting/5/present_user_ids": []}), 0)
        self.assertFalse(Anwesenheit.objects.exists())

    def test_fehlerantwort_loescht_nichts_und_wird_gemeldet(self):
        from apps.events.models import Anwesenheit
        self._abrufen(_anwesenheit_antwort())
        with self.assertRaises(OpenSlidesFehler):
            self._abrufen({"success": False, "message": "Kein Zugriff"})
        self.assertEqual(Anwesenheit.objects.count(), 2)

    def test_ohne_meeting_wird_abgelehnt(self):
        self.ver.openslides_meeting_id = None
        self.ver.save()
        with self.assertRaises(OpenSlidesFehler):
            services.anwesenheit_abrufen(self.verbindung, self.ver)

    def test_platzhalter_anwesenheit_und_anzahl(self):
        from apps.documents.platzhalter import kontext
        ctx = kontext(self.v, veranstaltung=self.ver)
        self.assertEqual(ctx["anwesenheit_anzahl"], "[Anzahl]")
        self.assertEqual(ctx["anwesenheit"], "")   # leer -> kein "keine Angaben"-Feld
        self._abrufen(_anwesenheit_antwort())
        ctx = kontext(self.v, veranstaltung=self.ver)
        self.assertEqual(ctx["anwesenheit_anzahl"], "2")
        self.assertEqual(ctx["anwesenheit"], "Anwesenheitsliste:\n- Erika Muster\n- Gast Besucher")
        self.assertEqual(kontext(self.v, veranstaltung=self.ver, anlagen=True)["anwesenheit"],
                         "Anwesenheitsliste: siehe Anlage 1 (Teilnehmerliste)")

    def test_anonymisiertes_mitglied_erscheint_anonymisiert(self):
        from apps.documents.platzhalter import kontext
        self._abrufen(_anwesenheit_antwort())
        self.m.vorname, self.m.nachname = "Anonymisiert", "#7"
        self.m.save()
        self.assertIn("Anonymisiert #7", kontext(self.v, veranstaltung=self.ver)["anwesenheit"])
        self.assertNotIn("Erika", kontext(self.v, veranstaltung=self.ver)["anwesenheit"])


class ProtokollAnwesenheitViewTests(ProtokollAbstimmungenViewTests):
    """Erbt Einrichtung und Hilfsmethode; prueft zusaetzlich die Teilnehmerliste/Verweise im Protokolltext."""

    def test_verweis_im_text_und_veranstaltung_zeigt_abschnitt(self):
        self._post(anwesenheit=_anwesenheit_antwort())
        self.s.text = "Anwesend: {anwesenheit_anzahl}\n\n{anwesenheit}"
        self.s.save()
        from apps.documents.platzhalter import ersetzen
        text = ersetzen(self.s.text, self.s.kontext())
        self.assertIn("Anwesend: 2", text)
        self.assertIn("Anwesenheitsliste: siehe Anlage 1 (Teilnehmerliste)", text)
        r = self.client.get(reverse("veranstaltung_detail", args=[self.ver.pk]))
        self.assertContains(r, "Anwesenheit (aus OpenSlides)")
        self.assertContains(r, "Gast Besucher")


class LoginFehlerHinweisTests(TestCase):
    """Fehlermeldung beim Anmelden: bei Proxy-Fehlern (leere Antwort) ein konkreter Suchhinweis."""

    def _login(self, status, text=""):
        from types import SimpleNamespace
        from .client import OSClient
        v = SimpleNamespace(url="https://os.example.org", tls_pruefen=True, benutzername="u", passwort="p")
        c = OSClient(v)
        antwort = SimpleNamespace(status_code=status, text=text, headers={}, json=lambda: {})
        with patch.object(c.s, "post", return_value=antwort):
            with self.assertRaises(OpenSlidesFehler) as ctx:
                c.login()
        return str(ctx.exception)

    def test_502_nennt_ursache_und_befehle_statt_leerer_antwort(self):
        meldung = self._login(502)
        self.assertIn("HTTP 502).", meldung)
        self.assertIn("docker compose ps", meldung)
        self.assertNotIn("): ", meldung)   # keine leere Antwort hinter dem Doppelpunkt

    def test_antworttext_wird_weiter_angezeigt(self):
        self.assertIn("HTTP 400): kaputt", self._login(400, "kaputt"))

    def test_zugangsdaten_und_adresse_hinweise(self):
        self.assertIn("technischen Benutzers", self._login(403, "verboten"))
        self.assertIn("OpenSlides-Startseite", self._login(404, "nicht da"))


def _wahl_antwort_mit_benutzern():
    """Echte Form einer OpenSlides-4.3-Wahl: die Optionen verweisen direkt auf Benutzer (user/ID, keine
    Kandidatenliste) und die Wahlmethode ist "Y" (nur Stimmen) - Optionen absichtlich NICHT nach Stimmen sortiert."""
    return {
        "meeting/5/assignment_ids": [1],
        "assignment/1/title": "wahl vorsitzender",
        "assignment/1/poll_ids": [3],
        "poll/3/title": "Wahlgang", "poll/3/state": "published", "poll/3/pollmethod": "Y",
        "poll/3/option_ids": [14, 13, 15],
        "option/13/yes": "10.000000", "option/13/content_object_id": "user/21",
        "option/14/yes": "5.000000", "option/14/content_object_id": "user/22",
        "option/15/yes": "2.000000", "option/15/content_object_id": "user/23",
        "user/21/first_name": "Andreas", "user/21/last_name": "Hoffmann",
        "user/22/first_name": "Karin", "user/22/last_name": "Hofmann",
        "user/23/first_name": "Patrick", "user/23/last_name": "Meier",
    }


class WahlMitBenutzerOptionenTests(TestCase):
    """Regression: im Protokoll stand "Option 13: 10 Ja" statt "Andreas Hoffmann: 10 Stimmen"."""

    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)

    def _abrufen(self, antwort):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.return_value = antwort
            return services.wahlergebnisse_abrufen(self.verbindung, self.ver)

    def test_namen_und_stimmen_statt_option_nummern(self):
        self.assertEqual(self._abrufen(_wahl_antwort_mit_benutzern()), 1)
        w = Wahlergebnis.objects.get(veranstaltung=self.ver)
        self.assertEqual(w.amt, "wahl vorsitzender")
        self.assertEqual(w.wahlgang, "Wahlgang")
        self.assertEqual(w.ergebnis.splitlines(), ["Andreas Hoffmann: 10 Stimmen", "Karin Hofmann: 5 Stimmen",
                                                   "Patrick Meier: 2 Stimmen"])
        self.assertNotIn("Option", w.ergebnis)
        self.assertNotIn("Ja", w.ergebnis)

    def test_anfrage_holt_benutzernamen_der_optionen(self):
        felder = services._wahlergebnisse_anfrage(5)[0]["fields"]["assignment_ids"]["fields"]["poll_ids"]["fields"][
            "option_ids"]["fields"]["content_object_id"]["fields"]
        for feld in ("first_name", "last_name", "username", "poll_candidate_ids"):
            self.assertIn(feld, felder)

    def test_ja_nein_enthaltung_bleibt_bei_anderen_wahlmethoden(self):
        antwort = _wahl_antwort_mit_benutzern()
        antwort["poll/3/pollmethod"] = "YN"
        antwort["option/13/no"] = "1.000000"
        self._abrufen(antwort)
        ergebnis = Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis
        self.assertIn("Andreas Hoffmann: 10 Ja, 1 Nein", ergebnis)

    def test_ohne_aufloesbaren_namen_bleibt_der_notnamen(self):
        antwort = _wahl_antwort_mit_benutzern()
        for k in [k for k in antwort if k.startswith("user/21")]:
            del antwort[k]
        self._abrufen(antwort)
        self.assertIn("Nutzer 21: 10 Stimmen", Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis)

    def test_protokolltext_zeigt_die_wahl_lesbar(self):
        from apps.documents.platzhalter import kontext
        self._abrufen(_wahl_antwort_mit_benutzern())
        text = kontext(self.v, veranstaltung=self.ver)["wahlergebnisse"]
        self.assertIn("wahl vorsitzender – Wahlgang:\nAndreas Hoffmann: 10 Stimmen", text)


def _antrag_mit_details():
    """Antrag wie in OpenSlides 4.3 (Antragsteller ueber meeting_user, HTML-Text/Begruendung, Zaehlwerte -2 =
    "nicht erfasst")."""
    return {
        "meeting/5/motion_ids": [2],
        "motion/2/title": "beitragserhöhung", "motion/2/state_id": 4, "motion_state/4/name": "angenommen",
        "motion/2/text": "<p>erhöhung beitrag</p>", "motion/2/reason": "<p>kein geld in der <strong>kasse</strong></p>",
        "motion/2/submitter_ids": [8], "motion_submitter/8/meeting_user_id": 31,
        "meeting_user/31/user_id": 55, "user/55/first_name": "Christina", "user/55/last_name": "Peters",
        "motion/2/poll_ids": [9],
        "poll/9/title": "Antrag Abstimmung", "poll/9/state": "published", "poll/9/option_ids": [500],
        "poll/9/votesvalid": "-2.000000", "poll/9/votesinvalid": "-2.000000", "poll/9/votescast": "-2.000000",
        "option/500/yes": "10.000000", "option/500/no": "5.000000", "option/500/abstain": "2.000000",
    }


class AntragDetailsTests(TestCase):
    """Regression: im Protokoll fehlten Antragsteller, Antragstext und Begruendung, und -2 wurde als Stimmenzahl
    ausgegeben."""

    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)

    def _abrufen(self, antraege):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = [{}, antraege, {}]
            return services.wahlergebnisse_abrufen(self.verbindung, self.ver)

    def test_antragsteller_text_begruendung_ergebnis_und_status(self):
        self.assertEqual(self._abrufen(_antrag_mit_details()), 1)
        a = Wahlergebnis.objects.get(veranstaltung=self.ver)
        self.assertEqual(a.amt, "Antrag: beitragserhöhung")
        self.assertEqual(a.ergebnis.splitlines(), [
            "Antragsteller*in: Christina Peters", "Antragstext: erhöhung beitrag", "Begründung: kein geld in der kasse",
            "Ergebnis: 10 Ja, 5 Nein, 2 Enthaltung", "Status in OpenSlides: angenommen"])

    def test_nicht_erfasste_zaehlwerte_minus_zwei_werden_ausgeblendet(self):
        self._abrufen(_antrag_mit_details())
        ergebnis = Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis
        self.assertNotIn("-2", ergebnis)
        self.assertNotIn("Gültige", ergebnis)

    def test_erfasste_zaehlwerte_werden_gezeigt(self):
        antwort = _antrag_mit_details()
        antwort.update({"poll/9/votesvalid": "17.000000", "poll/9/votesinvalid": "0.000000",
                        "poll/9/votescast": "17.000000"})
        self._abrufen(antwort)
        self.assertIn("Gültige Stimmen: 17, Ungültige Stimmen: 0, Abgegebene Stimmen: 17",
                      Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis)

    def test_antragsteller_direkt_ueber_user_aeltere_version_und_mehrere(self):
        antwort = _antrag_mit_details()
        antwort["motion/2/submitter_ids"] = [8, 9]
        del antwort["motion_submitter/8/meeting_user_id"]
        antwort["motion_submitter/8/user_id"] = 55
        antwort["motion_submitter/9/user_id"] = 56
        antwort["user/56/first_name"], antwort["user/56/last_name"] = "Karin", "Hofmann"
        self._abrufen(antwort)
        self.assertIn("Antragsteller*in: Christina Peters, Karin Hofmann",
                      Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis)

    def test_ohne_antragsteller_text_und_begruendung_keine_leeren_zeilen(self):
        antwort = {k: v for k, v in _antrag_mit_details().items()
                   if not k.startswith(("motion/2/text", "motion/2/reason", "motion/2/submitter"))}
        self._abrufen(antwort)
        ergebnis = Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis
        self.assertNotIn("Antragsteller", ergebnis)
        self.assertNotIn("Antragstext", ergebnis)
        self.assertTrue(ergebnis.startswith("Ergebnis: 10 Ja"))

    def test_anfrage_holt_text_begruendung_und_antragsteller(self):
        felder = services._antraege_anfrage(5)[0]["fields"]["motion_ids"]["fields"]
        for feld in ("text", "reason", "submitter_ids", "number", "state_id"):
            self.assertIn(feld, felder)
        self.assertIn("meeting_user_id", felder["submitter_ids"]["fields"])
        alt = services._antraege_anfrage(5, "user")[0]["fields"]["motion_ids"]["fields"]
        self.assertIn("user_id", alt["submitter_ids"]["fields"])
        ohne = services._antraege_anfrage(5, "ohne_antragsteller")[0]["fields"]["motion_ids"]["fields"]
        self.assertNotIn("submitter_ids", ohne)
        self.assertIn("text", ohne)
        minimal = services._antraege_anfrage(5, "minimal")[0]["fields"]["motion_ids"]["fields"]
        self.assertEqual(set(minimal), {"number", "title", "state_id", "poll_ids"})

    def test_klartext_aus_html(self):
        self.assertEqual(services._klartext("<p>Zeile 1</p><p>Zeile&nbsp;2 &amp; mehr</p>"), "Zeile 1\nZeile\xa02 & mehr")
        self.assertEqual(services._klartext("a<br>b<br/>c"), "a\nb\nc")
        self.assertEqual(services._klartext(None), "")
        self.assertTrue(services._klartext("x" * 6000).endswith("[...]"))

    def test_thema_text_wird_uebernommen(self):
        thema = {"meeting/5/topic_ids": [9], "topic/9/title": "Entlastung", "topic/9/text": "<p>Der Vorstand wird entlastet.</p>",
                 "topic/9/poll_ids": [70], "poll/70/title": "Abstimmung", "poll/70/state": "finished",
                 "poll/70/option_ids": [400], "option/400/yes": "10", "option/400/no": "0", "option/400/abstain": "2"}
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = [{}, {}, thema]
            services.wahlergebnisse_abrufen(self.verbindung, self.ver)
        self.assertTrue(Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis.startswith(
            "Text: Der Vorstand wird entlastet."))


class AbfrageRueckfallTests(TestCase):
    """Lehnt OpenSlides eine Abfrage-Variante ab (Feld gibt es in dieser Version nicht), muss es mit weniger Feldern
    weitergehen - die Abstimmungen duerfen nie stillschweigend verschwinden."""

    def setUp(self):
        self.v = Verein.objects.create(name="Test e.V.", kuerzel="test")
        self.verbindung = OpenSlidesVerbindung.objects.create(
            verein=self.v, url="https://os.example.org", benutzername="sync", passwort="geheim")
        self.ver = Veranstaltung.objects.create(verein=self.v, titel="Mitgliederversammlung 2026",
                                                beginn=timezone.now(), openslides_meeting_id=5)

    def _abrufen(self, *antworten):
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = list(antworten)
            return services.wahlergebnisse_abrufen(self.verbindung, self.ver), MockClient.return_value

    def test_abgelehnte_erste_variante_faellt_auf_die_naechste_zurueck(self):
        antwort = _antrag_mit_details()
        antwort["motion_submitter/8/user_id"] = 55   # aeltere Version: Antragsteller direkt ueber user
        del antwort["motion_submitter/8/meeting_user_id"]
        n, client = self._abrufen({}, OpenSlidesFehler("Abfrage abgelehnt (HTTP 400): unbekanntes Feld"), antwort, {})
        self.assertEqual(n, 1)
        self.assertIn("Antragsteller*in: Christina Peters", Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis)
        self.assertEqual(len(self.ver.abruf_hinweise), 1)
        self.assertIn('Variante "user"', self.ver.abruf_hinweise[0])

    def test_fehlerantwort_dict_zaehlt_wie_ein_fehler(self):
        n, _ = self._abrufen({}, {"message": "unbekanntes Feld"}, _antrag_mit_details(), {})
        self.assertEqual(n, 1)   # zweite Variante liefert die Daten

    def test_minimale_variante_liefert_wenigstens_die_abstimmung(self):
        minimal = {k: v for k, v in _antrag_mit_details().items()
                   if not k.startswith(("motion/2/text", "motion/2/reason", "motion/2/submitter", "motion_submitter",
                                        "meeting_user", "user/"))}
        fehler = OpenSlidesFehler("HTTP 400")
        n, _ = self._abrufen({}, fehler, fehler, fehler, minimal, {})
        self.assertEqual(n, 1)
        ergebnis = Wahlergebnis.objects.get(veranstaltung=self.ver).ergebnis
        self.assertIn("Ergebnis: 10 Ja, 5 Nein, 2 Enthaltung", ergebnis)
        self.assertNotIn("Antragsteller", ergebnis)
        self.assertIn('Variante "minimal"', self.ver.abruf_hinweise[0])

    def test_vollstaendige_variante_ohne_hinweis(self):
        self._abrufen({}, _antrag_mit_details(), {})
        self.assertEqual(self.ver.abruf_hinweise, [])

    def test_themen_ohne_text_variante(self):
        thema = {"meeting/5/topic_ids": [9], "topic/9/title": "Entlastung", "topic/9/poll_ids": [70],
                 "poll/70/title": "Abstimmung", "poll/70/state": "finished", "poll/70/option_ids": [400],
                 "option/400/yes": "10", "option/400/no": "0", "option/400/abstain": "2"}
        n, _ = self._abrufen({}, {}, OpenSlidesFehler("HTTP 400"), thema)
        self.assertEqual(n, 1)
        self.assertEqual(Wahlergebnis.objects.get().amt, "Thema: Entlastung")

    def test_details_sind_als_paare_gespeichert(self):
        self._abrufen({}, _antrag_mit_details(), {})
        paare = Wahlergebnis.objects.get(veranstaltung=self.ver).details
        self.assertIn(["Antragsteller*in", "Christina Peters"], paare)
        self.assertIn(["Ergebnis", "10 Ja, 5 Nein, 2 Enthaltung"], paare)
        self.assertIn(["Status in OpenSlides", "angenommen"], paare)

    def test_wahl_details_name_und_stimmen(self):
        self._abrufen(_wahl_antwort_mit_benutzern(), {}, {})
        self.assertEqual(Wahlergebnis.objects.get(art="wahl").details,
                         [["Andreas Hoffmann", "10 Stimmen"], ["Karin Hofmann", "5 Stimmen"],
                          ["Patrick Meier", "2 Stimmen"]])

    def test_client_meldet_http_fehler_der_abfrage(self):
        from types import SimpleNamespace
        from .client import OSClient
        c = OSClient(SimpleNamespace(url="https://os.example.org", tls_pruefen=True, benutzername="u", passwort="p"))
        c.token = "t"
        antwort = SimpleNamespace(status_code=400, text="", json=lambda: {"message": "unknown field"})
        with patch.object(c.s, "post", return_value=antwort):
            with self.assertRaisesMessage(OpenSlidesFehler, "Abfrage abgelehnt (HTTP 400)"):
                c.abfragen([{"ids": [1], "collection": "meeting", "fields": {}}])

    def test_ansicht_zeigt_hinweis_bei_eingeschraenktem_abruf(self):
        User = get_user_model()
        User.objects.create_superuser("admin", password="pw-Test-12345")
        Zugang.objects.create(verein=self.v, user=User.objects.get(username="admin"),
                              rolle=Rolle.objects.get(verein=self.v, name="Superadministrator"))
        self.client.login(username="admin", password="pw-Test-12345")
        self.verbindung.aktiv = True
        self.verbindung.save()
        with patch("apps.openslides.services.OSClient") as MockClient:
            MockClient.return_value.abfragen.side_effect = [{}, OpenSlidesFehler("HTTP 400"), _antrag_mit_details(),
                                                            {}, {}]
            r = self.client.post(reverse("veranstaltung_wahlergebnisse", args=[self.ver.pk]), follow=True)
        self.assertContains(r, "nur eingeschraenkt abgerufen")
