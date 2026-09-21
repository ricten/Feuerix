"""Word-Export (.docx) - damit Schriftstücke außerhalb des Systems weiterbearbeitet werden können.
Briefkopf/Fuß folgen dem gleichen CI wie die PDF-Erzeugung (apps/core/pdf.py) und der Original-Vorlage:
großer Vereinsname in neutralem Grau, Logo oben rechts frei positioniert (keine Tabelle - die erzeugt
in Word beim Bearbeiten störende Rahmen-/Gitternetzlinien), Trennlinie in der Akzentfarbe, Fußzeile mit
Vertretungsberechtigtem und Bankverbindung."""
from io import BytesIO

from .pdf import absaetze
from .platzhalter import ersetzen

STANDARD_AKZENTFARBE = "1F4E79"
BRIEFKOPF_TEXTFARBE = "646363"
LOGO_BREITE_CM = 5.5
LOGO_HOEHE_CM = 2.8


def _farbe(verein):
    from docx.shared import RGBColor
    hex_ = (verein.akzentfarbe or "").lstrip("#")
    try:
        return RGBColor.from_string(hex_)
    except ValueError:
        return RGBColor.from_string(STANDARD_AKZENTFARBE)


def _farbe_fuss(verein):
    """Zweite Akzentfarbe für die Linie über der Fußzeile; ohne eigene Farbe wie oben."""
    from docx.shared import RGBColor
    hex_ = (verein.akzentfarbe_fuss or "").lstrip("#")
    try:
        return RGBColor.from_string(hex_)
    except ValueError:
        return _farbe(verein)


def _trennlinie(paragraph, farbe):
    """Fügt dem Absatz einen unteren Rahmen hinzu (python-docx kennt keine Linie/HR direkt)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pbdr = OxmlElement("w:pBdr")
    unten = OxmlElement("w:bottom")
    unten.set(qn("w:val"), "single")
    unten.set(qn("w:sz"), "18")
    unten.set(qn("w:space"), "1")
    unten.set(qn("w:color"), "%02X%02X%02X" % (farbe[0], farbe[1], farbe[2]))
    pbdr.append(unten)
    paragraph._p.get_or_add_pPr().append(pbdr)


def _logo_masse(pfad):
    """-> (Breite, Höhe) als docx.shared.Cm, seitenverhältnistreu in die CI-Logobox eingepasst."""
    from docx.shared import Cm
    try:
        from PIL import Image
        with Image.open(pfad) as bild:
            iw, ih = bild.size
        skala = min(LOGO_BREITE_CM / iw, LOGO_HOEHE_CM / ih)
        return Cm(iw * skala), Cm(ih * skala)
    except Exception:
        return Cm(LOGO_BREITE_CM), Cm(LOGO_HOEHE_CM)


def _als_schwebend(bild_shape):
    """Wandelt ein inline eingefügtes Bild in ein frei positioniertes Objekt oben rechts am Seitenrand um -
    wie im Original-Briefkopf, der Logo und Trennlinie ebenfalls frei positioniert statt in einer Tabelle
    platziert (eine Tabelle zeigt beim Bearbeiten in Word sonst störende Rahmenlinien)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    inline = bild_shape._inline
    drawing = inline.getparent()
    anchor = OxmlElement("wp:anchor")
    for attr, wert in (("distT", "0"), ("distB", "0"), ("distL", "114300"), ("distR", "114300"),
                       ("simplePos", "0"), ("relativeHeight", "251658240"), ("behindDoc", "0"),
                       ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1")):
        anchor.set(attr, wert)
    simplepos = OxmlElement("wp:simplePos")
    simplepos.set("x", "0")
    simplepos.set("y", "0")
    anchor.append(simplepos)
    posh = OxmlElement("wp:positionH")
    posh.set("relativeFrom", "margin")
    ah = OxmlElement("wp:align")
    ah.text = "right"
    posh.append(ah)
    anchor.append(posh)
    posv = OxmlElement("wp:positionV")
    posv.set("relativeFrom", "margin")
    av = OxmlElement("wp:align")
    av.text = "top"
    posv.append(av)
    anchor.append(posv)
    for tag in ("wp:extent", "wp:effectExtent"):
        el = inline.find(qn(tag))
        if el is not None:
            anchor.append(el)
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        el = inline.find(qn(tag))
        if el is not None:
            anchor.append(el)
    drawing.replace(inline, anchor)


def _fusszeile(doc, verein):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    linie = doc.sections[0].footer.paragraphs[0]
    linie.paragraph_format.space_after = Pt(6)
    _trennlinie(linie, _farbe_fuss(verein))
    absatz = doc.sections[0].footer.add_paragraph()
    absatz.alignment = WD_ALIGN_PARAGRAPH.CENTER
    zeilen = [z for z in (
        " | ".join(x for x in (verein.name, verein.vereinsregister, verein.email) if x),
        ", ".join(x for x in (verein.unterschrift_1, verein.adresszeile) if x),
        " | ".join(x for x in (
            f"Bank: {verein.bankname}" if verein.bankname else "",
            f"IBAN: {verein.iban}" if verein.iban else "",
            f"BIC: {verein.bic}" if verein.bic else "") if x),
    ) if z]
    for i, z in enumerate(zeilen):
        lauf = absatz.add_run(z)
        lauf.font.size = Pt(10)
        lauf.font.color.rgb = RGBColor.from_string(BRIEFKOPF_TEXTFARBE)
        if i < len(zeilen) - 1:
            lauf.add_break()


def schriftstueck_docx(s):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    ctx = s.kontext()
    v = s.verein
    farbe = _farbe(v)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)

    # Briefkopf: Vereinsname (groß, neutrales Grau), Logo frei oben rechts positioniert (keine Tabelle,
    # sonst zeigt Word beim Bearbeiten Rahmenlinien); die Akzentfarbe wird nur für die Trennlinien verwendet.
    logo_pfad = None
    try:
        if v.logo:
            logo_pfad = v.logo.path
    except Exception:
        pass
    logo_breite = logo_hoehe = None
    if logo_pfad:
        logo_breite, logo_hoehe = _logo_masse(logo_pfad)

    kopf = doc.add_paragraph()
    if logo_pfad:
        kopf.paragraph_format.right_indent = logo_breite + Cm(0.4)
    lauf = kopf.add_run(v.name)
    lauf.font.size = Pt(24)
    lauf.font.color.rgb = RGBColor.from_string(BRIEFKOPF_TEXTFARBE)
    lauf.font.name = "Montserrat"
    lauf._element.rPr.rFonts.set(qn("w:hAnsi"), "Montserrat")
    if logo_pfad:
        try:
            bild = kopf.add_run().add_picture(logo_pfad, width=logo_breite, height=logo_hoehe)
            _als_schwebend(bild)
        except Exception:
            pass

    linie = doc.add_paragraph()
    linie.paragraph_format.space_before = Pt(2)
    linie.paragraph_format.space_after = Pt(10)
    if logo_pfad:
        linie.paragraph_format.right_indent = logo_breite + Cm(0.4)
    _trennlinie(linie, farbe)

    absender = ", ".join(x for x in (v.name, v.anschrift, f"{v.plz} {v.ort}".strip()) if x)
    p = doc.add_paragraph(absender)
    p.runs[0].font.size = Pt(8)
    if s.mitglied_id and s.art != "protokoll":
        doc.add_paragraph("\n".join(z for z in s.mitglied.anschrift_zeilen() if z))
    d = doc.add_paragraph(ctx["datum"])
    d.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_heading(ersetzen(s.betreff or s.titel, ctx), level=2)
    for a in absaetze(ersetzen(s.text, ctx)):
        doc.add_paragraph(a)

    _fusszeile(doc, v)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
