"""Formatierter Text (aus dem Editor TinyMCE) fuer Vorlagen, Schriftstuecke und Serienbriefe.

Gespeichert wird bereinigtes HTML mit einer kleinen, festen Auswahl an Formatierungen (Absaetze, Ueberschriften, fett/kursiv/
unterstrichen, Listen, Tabellen, Ausrichtung, Links). Aeltere Texte (reiner Text mit Leerzeilen) bleiben gueltig: ist_html()
erkennt sie, aus_text() wandelt sie bei Bedarf um. Platzhalter wie {vorname} werden im HTML mit maskierten Werten ersetzt.
Aus dem HTML entstehen PDF-Bausteine (ReportLab), Word-Absaetze (python-docx) und reiner Text (E-Mail)."""
import re
from html import escape
from html.parser import HTMLParser

ERLAUBT = {"p", "br", "strong", "em", "u", "s", "h1", "h2", "h3", "ul", "ol", "li", "table", "thead", "tbody", "tfoot",
           "tr", "th", "td", "a"}
ERSATZ = {"b": "strong", "i": "em", "strike": "s", "del": "s", "div": "p", "h4": "h3", "h5": "h3", "h6": "h3"}
LEER = {"br"}
VERWERFEN = {"script", "style", "iframe", "object", "embed", "template"}   # samt Inhalt
AUSRICHTUNG = {"left", "right", "center", "justify"}
BLOCKTAGS = {"p", "h1", "h2", "h3", "ul", "ol", "table"}
_ALS_HTML = re.compile(r"^\s*<(p|h[1-6]|ul|ol|table|div)[\s>/]", re.I)
_PLATZHALTER = re.compile(r"\{([a-z_0-9]+)\}")
_SICHERE_LINKS = ("http://", "https://", "mailto:", "tel:")


class Knoten:
    __slots__ = ("tag", "attr", "kinder", "text")

    def __init__(self, tag=None, attr=None, text=None):
        self.tag, self.attr, self.kinder, self.text = tag, attr or {}, [], text


class _Baum(HTMLParser):
    """Baut aus HTML einen bereinigten Baum (nur erlaubte Tags/Attribute); Unbekanntes wird entfernt, der Text bleibt."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.wurzel = Knoten("wurzel")
        self.stapel = [self.wurzel]
        self.verwerfen = 0

    def _attribute(self, tag, attrs):
        a = dict(attrs)
        sauber = {}
        stil = a.get("style") or ""
        m = re.search(r"text-align\s*:\s*([a-z]+)", stil, re.I)
        if m and m.group(1).lower() in AUSRICHTUNG and tag in ("p", "h1", "h2", "h3", "li", "td", "th"):
            sauber["align"] = m.group(1).lower()
        if tag == "a":
            href = (a.get("href") or "").strip()
            if href.lower().startswith(_SICHERE_LINKS):
                sauber["href"] = href
        if tag in ("td", "th"):
            for k in ("colspan", "rowspan"):
                if (a.get(k) or "").isdigit() and 1 <= int(a[k]) <= 20:
                    sauber[k] = a[k]
        return sauber

    def handle_starttag(self, tag, attrs):
        if tag in VERWERFEN:
            self.verwerfen += 1
            return
        if self.verwerfen:
            return
        tag = ERSATZ.get(tag, tag)
        if tag not in ERLAUBT:
            return
        k = Knoten(tag, self._attribute(tag, attrs))
        self.stapel[-1].kinder.append(k)
        if tag not in LEER:
            self.stapel.append(k)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in LEER and ERSATZ.get(tag, tag) in ERLAUBT and tag not in VERWERFEN:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if tag in VERWERFEN:
            self.verwerfen = max(0, self.verwerfen - 1)
            return
        if self.verwerfen:
            return
        tag = ERSATZ.get(tag, tag)
        if tag not in ERLAUBT or tag in LEER:
            return
        for i in range(len(self.stapel) - 1, 0, -1):
            if self.stapel[i].tag == tag:
                del self.stapel[i:]
                return

    def handle_data(self, data):
        if not self.verwerfen and data:
            self.stapel[-1].kinder.append(Knoten(text=data))


def baum(html):
    p = _Baum()
    p.feed(html or "")
    p.close()
    return p.wurzel


def _tag_auf(k):
    attr = ""
    if k.attr.get("align"):
        attr += f' style="text-align: {k.attr["align"]}"'
    if k.attr.get("href"):
        attr += f' href="{escape(k.attr["href"], quote=True)}"'
    for n in ("colspan", "rowspan"):
        if k.attr.get(n):
            attr += f' {n}="{k.attr[n]}"'
    return f"<{k.tag}{attr}>"


def _serialisieren(k):
    if k.text is not None:
        return escape(k.text, quote=False)
    if k.tag == "wurzel":
        return "".join(_serialisieren(c) for c in k.kinder)
    if k.tag in LEER:
        return f"<{k.tag}>"
    return _tag_auf(k) + "".join(_serialisieren(c) for c in k.kinder) + f"</{k.tag}>"


def bereinigen(html):
    """Sicheres HTML aus beliebigem HTML (Whitelist). Lose Textstuecke ausserhalb von Absaetzen werden zu Absaetzen."""
    w = baum(html)
    ausgabe, lose = [], []

    def lose_leeren():
        if lose and "".join(_serialisieren(x) for x in lose).strip():
            ausgabe.append("<p>" + "".join(_serialisieren(x) for x in lose).strip() + "</p>")
        lose.clear()

    for k in w.kinder:
        if k.tag in BLOCKTAGS:
            lose_leeren()
            ausgabe.append(_serialisieren(k))
        else:
            lose.append(k)
    lose_leeren()
    return "\n".join(ausgabe)


def ist_html(text):
    return bool(_ALS_HTML.match(text or ""))


def aus_text(text):
    """Reinen Text (Leerzeile = neuer Absatz, einfacher Umbruch = Zeilenumbruch) in HTML wandeln."""
    absaetze = [a.strip("\n") for a in re.split(r"\n[ \t]*\n", (text or "").replace("\r\n", "\n")) if a.strip()]
    return "\n".join("<p>" + escape(a, quote=False).replace("\n", "<br>") + "</p>" for a in absaetze)


def als_html(text):
    """Gespeicherten Text (HTML oder alter reiner Text) als bereinigtes HTML."""
    return bereinigen(text) if ist_html(text) else aus_text(text)


def ersetzen_html(html, ctx):
    """Platzhalter in HTML ersetzen; Werte werden maskiert, Zeilenumbrueche zu <br>. Unbekannte Platzhalter bleiben."""
    def ersatz(m):
        if m.group(1) not in ctx:
            return m.group(0)
        return escape(str(ctx[m.group(1)]), quote=False).replace("\r\n", "\n").replace("\n", "<br>")
    return _PLATZHALTER.sub(ersatz, html or "")


def text_ersetzt(text, ctx):
    """Gespeicherten Text (HTML oder alt) mit eingesetzten Platzhaltern als sauberes HTML."""
    return bereinigen(ersetzen_html(als_html(text), ctx))


# --------------------------------------------------------------------------- reiner Text
def _text_von(k, teile, liste=None):
    if k.text is not None:
        teile.append(k.text.replace("\xa0", " "))
    elif k.tag == "br":
        teile.append("\n")
    elif k.tag in ("ul", "ol"):
        for i, li in enumerate([c for c in k.kinder if c.tag == "li"], 1):
            teile.append(f"\n{i}. " if k.tag == "ol" else "\n- ")
            for c in li.kinder:
                _text_von(c, teile)
    elif k.tag == "table":
        for tr in _zeilen(k):
            teile.append("\n" + "\t".join(_inline_text(c) for c in tr))
        teile.append("\n")
    else:
        inner = []
        for c in k.kinder:
            _text_von(c, inner)
        s = "".join(inner)
        if k.tag in ("p", "h1", "h2", "h3"):
            teile.append("\n\n" + s.strip() + "\n\n")
        else:
            teile.append(s)


def _inline_text(k):
    t = []
    _text_von(k, t)
    return "".join(t).strip()


def zu_text(html):
    """HTML -> reiner Text (Absaetze durch Leerzeile) - fuer E-Mail-Text und Pruefungen."""
    teile = []
    for k in baum(html).kinder:
        _text_von(k, teile)
    text = "".join(teile)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def ist_leer(html):
    return not zu_text(html).strip()


def _zeilen(tabelle):
    """[[Zellknoten ...], ...] einer Tabelle (thead/tbody/tfoot werden aufgeloest)."""
    zeilen = []
    for k in tabelle.kinder:
        if k.tag == "tr":
            zeilen.append([c for c in k.kinder if c.tag in ("td", "th")])
        elif k.tag in ("thead", "tbody", "tfoot"):
            zeilen += [[c for c in tr.kinder if c.tag in ("td", "th")] for tr in k.kinder if tr.tag == "tr"]
    return [z for z in zeilen if z]


# --------------------------------------------------------------------------- PDF (ReportLab)
def _inline_markup(k):
    """Kinder eines Blocks -> ReportLab-Absatzmarkup (<b>, <i>, <u>, <br/>, <a>)."""
    if k.text is not None:
        return escape(k.text.replace("\xa0", " "), quote=False)
    if k.tag == "br":
        return "<br/>"
    inner = "".join(_inline_markup(c) for c in k.kinder if c.tag not in ("ul", "ol", "table"))
    return {"strong": f"<b>{inner}</b>", "em": f"<i>{inner}</i>", "u": f"<u>{inner}</u>",
            "s": f"<strike>{inner}</strike>",
            "a": (f'<a href="{escape(k.attr.get("href", ""), quote=True)}" color="#0645AD"><u>{inner}</u></a>'
                  if k.attr.get("href") else inner)}.get(k.tag, inner)


def _markup(kinder):
    return "".join(_inline_markup(c) for c in kinder)


def flowables(html, stile, layout, breite):
    """Bereinigtes HTML -> Liste von ReportLab-Flowables. `stile`: Stile aus core.pdf._stile(); `breite`: nutzbare
    Seitenbreite in Punkten (fuer Tabellen)."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Spacer, Table, TableStyle

    from apps.core.pdf import SCHRIFTEN
    normal = stile["normal"]
    ausr = {"left": 0, "center": 1, "right": 2, "justify": 4}
    groessen = {"h1": 1.7, "h2": 1.35, "h3": 1.15}

    def stil(tag, attr):
        basis = normal
        if tag in groessen:
            g = layout.schriftgroesse * groessen[tag]
            basis = ParagraphStyle(f"rt_{tag}", parent=normal, fontName=SCHRIFTEN[layout.schriftart], fontSize=g,
                                   leading=g * 1.25, spaceBefore=4, spaceAfter=2, keepWithNext=1)
        if attr.get("align"):
            basis = ParagraphStyle(f"rt_{tag}_{attr['align']}", parent=basis, alignment=ausr[attr["align"]])
        return basis

    def block(k):
        """Ein Blockknoten -> Flowables."""
        if k.tag in ("p", "h1", "h2", "h3"):
            m = _markup(k.kinder).strip()
            if not m.replace("<br/>", "").strip():
                return [Spacer(1, layout.schriftgroesse * 1.35)]      # leerer Absatz = Leerzeile
            teile = [Paragraph(m, stil(k.tag, k.attr))]
            return teile + [Spacer(1, 3 * mm)] if k.tag == "p" else teile
        if k.tag in ("ul", "ol"):
            return [liste(k), Spacer(1, 3 * mm)]
        if k.tag == "table":
            return tabelle(k) + [Spacer(1, 3 * mm)]
        return []

    def liste(k):
        punkte = []
        for li in [c for c in k.kinder if c.tag == "li"]:
            inhalt, lose = [], []
            for c in li.kinder:
                if c.tag in ("ul", "ol"):
                    if lose:
                        inhalt.append(Paragraph(_markup(lose), stil("p", li.attr)))
                        lose = []
                    inhalt.append(liste(c))
                else:
                    lose.append(c)
            if lose or not inhalt:
                inhalt.insert(0, Paragraph(_markup(lose) or "&nbsp;", stil("p", li.attr)))
            punkte.append(ListItem(inhalt, leftIndent=14))
        return ListFlowable(punkte, bulletType="1" if k.tag == "ol" else "bullet", start=None if k.tag == "ul" else 1,
                            bulletFontName=normal.fontName, bulletFontSize=normal.fontSize, leftIndent=14,
                            bulletDedent=12 if k.tag == "ul" else 14)

    def tabelle(k):
        zeilen = _zeilen(k)
        if not zeilen:
            return []
        spalten = max(sum(int(c.attr.get("colspan", 1)) for c in z) for z in zeilen)
        zell_stil = ParagraphStyle("rt_zelle", parent=normal, fontSize=max(8, layout.schriftgroesse - 1),
                                   leading=max(8, layout.schriftgroesse - 1) * 1.25)
        daten, spans, kopf_zeilen = [], [], []
        for r, z in enumerate(zeilen):
            zeile, c_idx = [], 0
            for c in z:
                m = _markup(c.kinder).strip()
                if c.tag == "th":
                    m = f"<b>{m}</b>"
                st = zell_stil if not c.attr.get("align") else ParagraphStyle(
                    f"rt_z_{c.attr['align']}", parent=zell_stil, alignment=ausr[c.attr["align"]])
                zeile.append(Paragraph(m or "&nbsp;", st))
                cs, rs = int(c.attr.get("colspan", 1)), int(c.attr.get("rowspan", 1))
                if cs > 1 or rs > 1:
                    spans.append(("SPAN", (c_idx, r), (min(c_idx + cs - 1, spalten - 1), r + rs - 1)))
                zeile += [""] * (cs - 1)
                c_idx += cs
            zeile += [""] * (spalten - len(zeile))
            daten.append(zeile)
            if all(c.tag == "th" for c in z):
                kopf_zeilen.append(r)
        t = Table(daten, colWidths=[breite / spalten] * spalten, hAlign="LEFT", repeatRows=1 if kopf_zeilen == [0] else 0)
        stile_t = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                   ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)] + spans
        for r in kopf_zeilen:
            stile_t.append(("BACKGROUND", (0, r), (-1, r), colors.whitesmoke))
        t.setStyle(TableStyle(stile_t))
        return [t]

    ergebnis, lose = [], []

    def lose_leeren():
        if lose:
            m = _markup(lose).strip()
            if m:
                ergebnis.extend([Paragraph(m, normal), Spacer(1, 3 * mm)])
            lose.clear()

    for k in baum(html).kinder:
        if k.tag in BLOCKTAGS:
            lose_leeren()
            ergebnis.extend(block(k))
        else:
            lose.append(k)
    lose_leeren()
    return ergebnis


# --------------------------------------------------------------------------- Word (python-docx)
def in_docx(doc, html):
    """Haengt den formatierten Text als Absaetze/Listen/Tabellen an ein python-docx-Dokument an."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    ausr = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}

    def laeufe(absatz, kinder, fett=False, kursiv=False, unter=False, durch=False):
        for c in kinder:
            if c.text is not None:
                if c.text.replace("\xa0", " "):
                    r = absatz.add_run(c.text.replace("\xa0", " "))
                    r.bold, r.italic, r.underline, r.font.strike = fett or None, kursiv or None, unter or None, durch or None
            elif c.tag == "br":
                absatz.add_run().add_break()
            elif c.tag in ("ul", "ol", "table"):
                continue
            else:
                laeufe(absatz, c.kinder, fett or c.tag == "strong", kursiv or c.tag == "em", unter or c.tag in ("u", "a"),
                       durch or c.tag == "s")

    def liste(k, ebene=0):
        stilname = "List Number" if k.tag == "ol" else "List Bullet"
        for li in [c for c in k.kinder if c.tag == "li"]:
            a = doc.add_paragraph(style=stilname + (f" {ebene + 1}" if ebene else ""))
            laeufe(a, [c for c in li.kinder if c.tag not in ("ul", "ol")])
            for c in li.kinder:
                if c.tag in ("ul", "ol"):
                    liste(c, min(ebene + 1, 2))

    lose = []

    def lose_leeren():
        if lose:
            a = doc.add_paragraph()
            laeufe(a, lose)
            lose.clear()

    for k in baum(html).kinder:
        if k.tag in ("p",):
            lose_leeren()
            a = doc.add_paragraph()
            laeufe(a, k.kinder)
            if k.attr.get("align"):
                a.alignment = ausr[k.attr["align"]]
        elif k.tag in ("h1", "h2", "h3"):
            lose_leeren()
            h = doc.add_heading(level={"h1": 1, "h2": 2, "h3": 3}[k.tag])
            laeufe(h, k.kinder)
        elif k.tag in ("ul", "ol"):
            lose_leeren()
            liste(k)
        elif k.tag == "table":
            lose_leeren()
            zeilen = _zeilen(k)
            if zeilen:
                spalten = max(len(z) for z in zeilen)
                t = doc.add_table(rows=len(zeilen), cols=spalten)
                t.style = "Table Grid"
                for i, z in enumerate(zeilen):
                    for j, c in enumerate(z):
                        a = t.cell(i, j).paragraphs[0]
                        laeufe(a, c.kinder, fett=c.tag == "th")
        else:
            lose.append(k)
    lose_leeren()
