"""Formatierter Text (TinyMCE): Bereinigung, Platzhalter, PDF/Word/E-Mail und Formular."""
import json
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core import mail
from django.test import TestCase
from django.urls import reverse
from pypdf import PdfReader

from apps.core.models import Rolle, Verein, Zugang
from apps.members.models import Mitglied

from . import richtext
from .docx_export import schriftstueck_docx
from .models import Schriftstueck, Serienbrief, Vorlage
from .pdf import schriftstueck_pdf, serienbrief_pdf
from .tasks import serienbrief_mailen_task

PW = "pw-Test-12345"


def pdf_text(pdf):
    return "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)


class BereinigenTests(TestCase):
    def test_gefaehrliches_wird_entfernt(self):
        h = ('<p onclick="x()" style="color:red; text-align: center">Hallo<script>alert(1)</script> '
             '<a href="javascript:alert(1)">böse</a> <a href="https://example.org/x?a=1&b=2">gut</a></p>'
             '<iframe src="x"></iframe><img src=x onerror=y><style>p{}</style>')
        s = richtext.bereinigen(h)
        for verboten in ("script", "onclick", "javascript", "iframe", "<img", "onerror", "<style", "color:red"):
            self.assertNotIn(verboten, s)
        self.assertIn('style="text-align: center"', s)
        self.assertIn('href="https://example.org/x?a=1&amp;b=2"', s)
        self.assertIn("Hallo", s)
        self.assertIn("böse", s)       # Text bleibt, nur der unsichere Link entfaellt

    def test_tags_werden_vereinheitlicht_und_lose_texte_zu_absaetzen(self):
        s = richtext.bereinigen("lose <b>fett</b> und <i>kursiv</i><div>Block</div>")
        self.assertEqual(s, "<p>lose <strong>fett</strong> und <em>kursiv</em></p>\n<p>Block</p>")

    def test_listen_und_tabellen_bleiben(self):
        h = "<ul><li>a</li><li>b</li></ul><table><tbody><tr><th>K</th><td colspan='2'>W</td></tr></tbody></table>"
        s = richtext.bereinigen(h)
        self.assertIn("<ul><li>a</li><li>b</li></ul>", s)
        self.assertIn('<td colspan="2">W</td>', s)

    def test_klartext_wird_zu_html_und_erkannt(self):
        self.assertFalse(richtext.ist_html("Hallo\n\nWelt"))
        self.assertFalse(richtext.ist_html("<b>kein Block</b>"))
        self.assertTrue(richtext.ist_html("<p>Hallo</p>"))
        self.assertEqual(richtext.aus_text("A & B\nzweite Zeile\n\nNeuer Absatz <x>"),
                         "<p>A &amp; B<br>zweite Zeile</p>\n<p>Neuer Absatz &lt;x&gt;</p>")
        self.assertEqual(richtext.als_html("Nur Text"), "<p>Nur Text</p>")

    def test_platzhalter_werden_maskiert_ersetzt(self):
        h = richtext.ersetzen_html("<p>{name}, {zeilen} {unbekannt}</p>", {"name": "Müller & <Söhne>", "zeilen": "a\nb"})
        self.assertEqual(h, "<p>Müller &amp; &lt;Söhne&gt;, a<br>b {unbekannt}</p>")

    def test_zu_text_und_leer(self):
        self.assertEqual(richtext.zu_text("<p>Eins</p><p>Zwei<br>drei</p><ul><li>x</li><li>y</li></ul>"),
                         "Eins\n\nZwei\ndrei\n\n- x\n- y")
        self.assertTrue(richtext.ist_leer("<p><br></p><p>&nbsp;</p>"))
        self.assertFalse(richtext.ist_leer("<p>x</p>"))


class AusgabeTests(TestCase):
    def setUp(self):
        self.verein = Verein.objects.create(name="Beispielverein e.V.", kuerzel="test", email="info@example.org")
        self.m = Mitglied.objects.create(verein=self.verein, anrede="herr", vorname="Max", nachname="Müller & Söhne",
                                         status="aktiv", email="max@example.org")
        self.html = ("<h1>Überschrift Eins</h1><p>Hallo {briefanrede} <strong>fetter</strong> <em>schiefer</em> "
                     "<u>unterstrichener</u> Text.</p><ul><li>Punkt A</li><li>Punkt B</li></ul>"
                     "<ol><li>Erstens</li><li>Zweitens</li></ol>"
                     "<table><tr><th>Kopf1</th><th>Kopf2</th></tr><tr><td>Zelle1</td><td>Zelle2</td></tr></table>"
                     "<p style='text-align:right'>Rechts</p>")

    def schriftstueck(self, text=None, **kw):
        return Schriftstueck.objects.create(verein=self.verein, titel="T", betreff="B", text=text or self.html,
                                            mitglied=self.m, **kw)

    def test_pdf_enthaelt_formatierungen(self):
        pdf = schriftstueck_pdf(self.schriftstueck())
        t = pdf_text(pdf)
        for erwartet in ("Überschrift Eins", "Hallo Sehr geehrter Herr Müller & Söhne,", "fetter", "Punkt A", "Punkt B",
                         "Erstens", "Zweitens", "Kopf1", "Zelle2", "Rechts"):
            self.assertIn(erwartet, t)
        self.assertIn(b"Helvetica-Bold", pdf)
        self.assertIn(b"Helvetica-Oblique", pdf)

    def test_alter_klartext_wird_wie_bisher_gesetzt(self):
        t = pdf_text(schriftstueck_pdf(self.schriftstueck(text="{briefanrede}\n\nErster Absatz.\nZeile zwei.\n\nZweiter.")))
        self.assertIn("Sehr geehrter Herr Müller & Söhne,", t)
        self.assertIn("Erster Absatz.\nZeile zwei.", t)
        self.assertIn("Zweiter.", t)

    def test_serienbrief_pdf_je_mitglied(self):
        sb = Serienbrief.objects.create(verein=self.verein, titel="S", betreff="B", text="<p>{vorname}: <strong>Hallo</strong></p>")
        t = pdf_text(serienbrief_pdf(sb, [self.m]))
        self.assertIn("Max: Hallo", t)

    def test_word_export(self):
        from docx import Document
        s = self.schriftstueck()
        doc = Document(BytesIO(schriftstueck_docx(s)))
        texte = [p.text for p in doc.paragraphs]
        for erwartet in ("Überschrift Eins", "Punkt A", "Erstens", "Rechts"):
            self.assertIn(erwartet, texte)
        fett = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
        self.assertIn("fetter", fett)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "Zelle2")

    def test_mail_hat_text_und_html_teil(self):
        sb = Serienbrief.objects.create(verein=self.verein, titel="S", betreff="Betreff", versandart="mail",
                                        text="<p>Hallo <strong>{vorname}</strong></p><ul><li>eins</li></ul>")
        serienbrief_mailen_task(sb.pk)
        self.assertEqual(len(mail.outbox), 1)
        m = mail.outbox[0]
        self.assertEqual(m.body, "Hallo Max\n\n- eins")
        html, art = m.alternatives[0].content, m.alternatives[0].mimetype
        self.assertEqual(art, "text/html")
        self.assertIn("<strong>Max</strong>", html)
        self.assertEqual(m.attachments[0][0], "S.pdf")


class EditorFormularTests(TestCase):
    def setUp(self):
        self.verein = Verein.objects.create(name="Beispielverein e.V.", kuerzel="test")
        self.admin = get_user_model().objects.create_superuser("admin", password=PW)
        Zugang.objects.create(verein=self.verein, user=self.admin,
                              rolle=Rolle.objects.get(verein=self.verein, name="Superadministrator"))
        self.client.force_login(self.admin)

    def test_formular_laedt_editor_und_platzhalter(self):
        for name in ("vorlage_add", "schriftstueck_add", "serienbrief_add"):
            r = self.client.get(reverse(name))
            self.assertContains(r, 'class="tinymce-editor"', msg_prefix=name)
            self.assertContains(r, "vendor/tinymce/tinymce.min.js")
            self.assertContains(r, "tinymce_init.js")
            self.assertNotRegex(r.content.decode(), r'<textarea[^>]*name="text"[^>]*required')
        gruppen = json.loads(self.client.get(reverse("vorlage_add")).context["form"].fields["text"].widget.attrs["data-platzhalter"])
        namen = {k for _t, es in gruppen for k, _b in es}
        self.assertIn("vorname", namen)
        self.assertNotIn("wahlergebnisse", namen)          # ohne OpenSlides-Anbindung nicht angeboten

    def test_speichern_bereinigt_html(self):
        daten = {"titel": "T", "art": "brief", "betreff": "B", "datum": "2026-01-01", "status": "entwurf",
                 "text": '<p onclick="x()">Hallo <strong>Welt</strong></p><script>alert(1)</script>'}
        r = self.client.post(reverse("schriftstueck_add"), daten)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Schriftstueck.objects.get().text, "<p>Hallo <strong>Welt</strong></p>")

    def test_leerer_editor_inhalt_wird_abgelehnt(self):
        daten = {"titel": "T", "art": "brief", "betreff": "B", "datum": "2026-01-01", "status": "entwurf",
                 "text": "<p><br></p><p>&nbsp;</p>"}
        r = self.client.post(reverse("schriftstueck_add"), daten)
        self.assertEqual(r.status_code, 200)
        self.assertFalse(Schriftstueck.objects.exists())

    def test_alter_klartext_erscheint_im_editor_als_html(self):
        v = Vorlage.objects.create(verein=self.verein, name="Alt", text="Hallo & Willkommen\n\nZweiter Absatz")
        r = self.client.get(reverse("vorlage_edit", args=[v.pk]))
        self.assertContains(r, "&lt;p&gt;Hallo &amp;amp; Willkommen&lt;/p&gt;")
        self.assertContains(r, "&lt;p&gt;Zweiter Absatz&lt;/p&gt;")
        v.refresh_from_db()
        self.assertEqual(v.text, "Hallo & Willkommen\n\nZweiter Absatz")      # in der Datenbank bleibt er bis zum Speichern

    def test_vorlage_aus_standardvorlagen_bleibt_nutzbar(self):
        self.client.post(reverse("vorlagen_standard"))
        self.assertTrue(Vorlage.objects.exists())
        for v in Vorlage.objects.all():
            self.assertTrue(richtext.flowables(richtext.text_ersetzt(v.text, {}), *self._stile_layout(), 400) is not None)

    def _stile_layout(self):
        from apps.core.pdf import Layout, _stile
        lay = Layout()
        return _stile(lay), lay

    def test_detailseite_zeigt_kein_rohes_html(self):
        v = Vorlage.objects.create(verein=self.verein, name="X", text="<p>Geheimer <strong>Text</strong></p>")
        r = self.client.get(reverse("vorlage_detail", args=[v.pk]))
        self.assertNotContains(r, "&lt;strong&gt;")
        self.assertNotContains(r, "<strong>Text</strong>")
