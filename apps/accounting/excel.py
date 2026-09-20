from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from .services import berichtsdaten


def kassenbericht_xlsx(b):
    d = berichtsdaten(b)
    wb = Workbook()
    ws = wb.active
    ws.title = "Zusammenfassung"
    fett = Font(bold=True)
    ws.append([b.titel]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Zeitraum {b.von:%d.%m.%Y} – {b.bis:%d.%m.%Y}", "", "Kassenwart", b.kassenwart])
    ws.append([])
    ws.append(["Konto", "Anfangsbestand", "Einnahmen", "Ausgaben", "Endbestand"])
    for c in ws[4]:
        c.font = fett
    for k in d["konten"]:
        ws.append([k["name"], float(k["anfang"]), float(k["einnahmen"]), float(k["ausgaben"]), float(k["ende"])])
    ws.append(["Summe", float(d["anfang"]), float(d["einnahmen"]), float(d["ausgaben"]), float(d["ende"])])
    for c in ws[ws.max_row]:
        c.font = fett
    ws.append([])
    ws.append(["Sphäre", "Einnahmen", "Ausgaben", "Ergebnis"])
    for c in ws[ws.max_row]:
        c.font = fett
    for l, e, a, r in d["sphaeren"]:
        ws.append([l, float(e), float(a), float(r)])
    for spalte, breite in zip("ABCDE", (38, 16, 16, 16, 16)):
        ws.column_dimensions[spalte].width = breite

    for titel, gruppen, gesamt, vj in (("Einnahmen", d["einnahmen_gruppen"], d["einnahmen"], d["vorjahr"]["einnahmen"]),
                                       ("Ausgaben", d["ausgaben_gruppen"], d["ausgaben"], d["vorjahr"]["ausgaben"])):
        w = wb.create_sheet(titel)
        w.append(["Sphäre", "Kategorie", "Betrag", f"Vorjahr ({d['vorjahr']['von']:%Y})"])
        for c in w[1]:
            c.font = fett
        for g in gruppen:
            for n, s, v in g["zeilen"]:
                w.append([g["label"], n, float(s), float(v)])
        w.append(["", "Summe", float(gesamt), float(vj)])
        for c in w[w.max_row]:
            c.font = fett
        for spalte, breite in zip("ABCD", (34, 40, 16, 16)):
            w.column_dimensions[spalte].width = breite

    j = wb.create_sheet("Kassenbuch")
    j.append(["Datum", "Beleg", "Buchungstext", "Kategorie", "Konto", "Einnahme", "Ausgabe"])
    for c in j[1]:
        c.font = fett
    for x in d["buchungen"]:
        j.append([x.datum, x.belegnummer, x.text, x.kategorie.name, x.konto.name,
                  float(x.betrag) if x.typ == "einnahme" else None, float(x.betrag) if x.typ == "ausgabe" else None])
        j.cell(j.max_row, 1).number_format = "DD.MM.YYYY"
    for spalte, breite in zip("ABCDEFG", (12, 16, 50, 30, 16, 14, 14)):
        j.column_dimensions[spalte].width = breite
    j.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
