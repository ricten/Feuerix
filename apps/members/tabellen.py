"""Einlesen von CSV/Excel-Tabellen und Wertekonvertierung (reines Python, ohne Django)."""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

# interne Feldnamen -> mögliche Spaltenüberschriften (normalisiert: klein, ohne Leer-/Sonderzeichen)
SYNONYME = {
    "mitgliedsnummer": ["mitgliedsnummer", "mitgliedsnr", "mitgliednr", "mitgliednummer", "nr", "nummer"],
    "anrede": ["anrede"],
    "vorname": ["vorname", "rufname"],
    "nachname": ["nachname", "familienname", "name"],
    "geburtsdatum": ["geburtsdatum", "geburtstag", "geboren", "geb", "gebdatum"],
    "eintrittsdatum": ["eintrittsdatum", "eintritt", "eintrittam", "mitgliedseit", "beitritt", "beitrittsdatum"],
    "austrittsdatum": ["austrittsdatum", "austritt", "austrittam"],
    "status": ["status"],
    "mitgliedsart": ["mitgliedsart", "beitragsart", "art", "mitgliedschaft"],
    "familie": ["familie", "familienzugehoerigkeit", "familienzugehorigkeit"],
    "ist_familienzahler": ["familienzahler", "zahltfamilienbeitrag", "istfamilienzahler"],
    "individueller_beitrag": ["individuellerbeitrag", "individuellerjahresbeitrag", "sonderbeitrag"],
    "strasse": ["strasse", "straße", "anschrift", "adresse", "strassehausnr", "strassenr"],
    "plz": ["plz", "postleitzahl"],
    "ort": ["ort", "wohnort", "stadt"],
    "email": ["email", "emailadresse", "mail", "emailadresse"],
    "telefon": ["telefon", "tel", "festnetz"],
    "mobil": ["mobil", "handy", "mobiltelefon"],
    "zahlungsart": ["zahlungsart", "zahlweise"],
    "kontoinhaber": ["kontoinhaber", "kontoinhaberin"],
    "iban": ["iban"],
    "bic": ["bic", "swift"],
    "mandatsreferenz": ["mandatsreferenz", "sepamandat", "mandat"],
    "mandatsdatum": ["mandatsdatum", "datumdessepamandats", "mandatvom"],
    "abteilungen": ["abteilung", "abteilungen"],
    "vorstandsmitglied": ["vorstandsmitglied", "vorstand"],
    "alters_ehrenabteilung": ["altersundehrenabteilung", "altersehrenabteilung", "ehrenabteilung", "altersabteilung"],
    "einsatzabteilung_aktiv": ["aktivesmitglieddereinsatzabteilung", "einsatzabteilung", "aktiveseinsatzabteilung",
                               "einsatzabteilungaktiv"],
    "notizen": ["notizen", "notiz", "bemerkung", "bemerkungen", "anmerkung"],
}
SPALTEN_ANZEIGE = [  # Reihenfolge und Überschriften der Vorlage / des Exports
    ("mitgliedsnummer", "Mitgliedsnummer"), ("anrede", "Anrede"), ("vorname", "Vorname"), ("nachname", "Nachname"),
    ("geburtsdatum", "Geburtsdatum"), ("eintrittsdatum", "Eintrittsdatum"), ("austrittsdatum", "Austrittsdatum"),
    ("status", "Status"), ("mitgliedsart", "Mitgliedsart"), ("familie", "Familie"),
    ("ist_familienzahler", "Familienzahler"), ("individueller_beitrag", "Individueller Beitrag"),
    ("strasse", "Straße"), ("plz", "PLZ"), ("ort", "Ort"), ("email", "E-Mail"), ("telefon", "Telefon"),
    ("mobil", "Mobil"), ("zahlungsart", "Zahlungsart"), ("kontoinhaber", "Kontoinhaber"), ("iban", "IBAN"),
    ("bic", "BIC"), ("mandatsreferenz", "Mandatsreferenz"), ("mandatsdatum", "Mandatsdatum"),
    ("abteilungen", "Abteilungen"), ("vorstandsmitglied", "Vorstandsmitglied"),
    ("alters_ehrenabteilung", "Alters- und Ehrenabteilung"),
    ("einsatzabteilung_aktiv", "Aktives Mitglied der Einsatzabteilung"), ("notizen", "Notizen"),
]
_LOOKUP = {}
for _feld, _namen in SYNONYME.items():
    for _n in _namen:
        _LOOKUP.setdefault(_n, _feld)


def norm(s):
    s = str(s or "").strip().lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z0-9]", "", s)


def feld_fuer(ueberschrift):
    n = norm(ueberschrift)
    return _LOOKUP.get(n) or _LOOKUP.get(n.replace("ae", "a").replace("oe", "o").replace("ue", "u"))


def lesen(dateiname, inhalt):
    """-> (Kopfzeile, Zeilen[list]) aus .xlsx oder .csv. Die erste nicht leere Zeile ist die Kopfzeile."""
    name = (dateiname or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        zeilen = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        try:
            text = inhalt.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = inhalt.decode("cp1252")
        probe = text[:4000]
        trenner = ";" if probe.count(";") >= probe.count(",") and probe.count(";") > 0 else ("\t" if probe.count("\t") > probe.count(",") else ",")
        zeilen = [list(r) for r in csv.reader(io.StringIO(text), delimiter=trenner)]
    else:
        raise ValueError("Bitte eine Excel-Datei (.xlsx) oder CSV-Datei (.csv) hochladen.")
    zeilen = [z for z in zeilen if any(str(c).strip() for c in z if c is not None)]
    if not zeilen:
        raise ValueError("Die Datei enthält keine Daten.")
    return [str(c or "").strip() for c in zeilen[0]], zeilen[1:]


def spaltenzuordnung(kopf):
    """-> ({index: feld}, [nicht erkannte Überschriften])"""
    zuordnung, unbekannt = {}, []
    for i, k in enumerate(kopf):
        f = feld_fuer(k)
        if f and f not in zuordnung.values():
            zuordnung[i] = f
        elif k:
            unbekannt.append(k)
    return zuordnung, unbekannt


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def datum(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Datum nicht lesbar: „{s}“ (erwartet TT.MM.JJJJ)")


def zahl(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).replace("€", "").replace(" ", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Zahl nicht lesbar: „{v}“")


def ja(v):
    return text(v).lower() in ("ja", "j", "x", "1", "true", "wahr", "yes", "y", "✓")


def anrede(v):
    s = text(v).lower().rstrip(".")
    return {"herr": "herr", "hr": "herr", "herrn": "herr", "m": "herr", "frau": "frau", "fr": "frau", "w": "frau",
            "divers": "divers", "d": "divers", "firma": "firma", "verein": "firma", "organisation": "firma"}.get(s, "")


def status(v):
    s = text(v).lower()
    return {"aktiv": "aktiv", "ruhend": "ruhend", "passiv": "ruhend", "ausgetreten": "ausgetreten", "gekuendigt": "ausgetreten",
            "gekündigt": "ausgetreten", "verstorben": "verstorben"}.get(s, "aktiv" if not s else None)


def zahlungsart(v):
    s = text(v).lower()
    if "lastschrift" in s or "sepa" in s:
        return "lastschrift"
    if s == "bar":
        return "bar"
    return "ueberweisung"


def iban_gueltig(iban):
    iban = re.sub(r"\s+", "", iban or "").upper()
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", iban):
        return False
    umgestellt = iban[4:] + iban[:4]
    zahlenstr = "".join(str(int(c, 36)) for c in umgestellt)
    return int(zahlenstr) % 97 == 1
