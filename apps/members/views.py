import json
from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.views.decorators.http import require_POST

from apps.core.crud import abschnitt, knopf, wert
from apps.core.models import AuditLog, Zugang

from .models import Mitglied


def mitglied_kontext(request, m):
    from apps.allowances.models import Aufwandsentschaedigung
    from apps.donations.models import Spende
    from apps.finance.models import Rechnung
    from apps.honors.models import Ehrung
    from apps.inventory.models import Verleih

    r = request.rechte
    aktionen, abschnitte = [], []
    if r.darf("mitglieder", "view"):
        aktionen.append(knopf("Datenauskunft (JSON)", reverse("mitglied_export", args=[m.pk])))
    if r.darf("selbstdienst", "change"):
        if not m.benutzer_id or not m.benutzer.is_active:
            aktionen.append(knopf("Zugangsdaten für Selbstdatenpflege senden", reverse("mitglied_zugang_einrichten", args=[m.pk]),
                                  post=True, stil="outline-primary",
                                  bestaetigung=f"Zugangsdaten per E-Mail an {m.email or '(keine E-Mail-Adresse hinterlegt!)'} senden?"))
        else:
            aktionen.append(knopf("Zugang zur Selbstdatenpflege sperren", reverse("mitglied_zugang_sperren", args=[m.pk]),
                                  post=True, stil="outline-danger", bestaetigung="Zugang sperren?"))
        if m.selbstdienst_initialpasswort:
            aktionen.append(knopf("Gespeichertes Startpasswort löschen", reverse("mitglied_startpasswort_loeschen", args=[m.pk]),
                                  post=True))
    verwaltungszugang = (Zugang.objects.filter(verein=request.verein, user_id=m.benutzer_id).select_related("rolle").first()
                        if m.benutzer_id else None)
    if verwaltungszugang and r.darf("verwaltung", "change"):
        status = "" if verwaltungszugang.aktiv else ", gesperrt"
        aktionen.append(knopf(f"Verwaltungszugang ({verwaltungszugang.rolle}{status}) bearbeiten",
                              reverse("zugang_edit", args=[verwaltungszugang.pk])))
    elif not verwaltungszugang and r.darf("verwaltung", "add"):
        aktionen.append(knopf("Verwaltungszugang einrichten", reverse("mitglied_verwaltungszugang_einrichten", args=[m.pk]),
                              stil="outline-primary"))
    if r.darf("mitglieder", "delete") and m.status != "verstorben" and m.vorname != "Anonymisiert":
        aktionen.append(knopf("Anonymisieren (DSGVO)", reverse("mitglied_anonymisieren", args=[m.pk]), post=True,
                              stil="outline-danger",
                              bestaetigung="Persönliche Daten unwiderruflich löschen? Rechnungen bleiben aus "
                                           "steuerlichen Gründen erhalten."))
    if r.darf("beitraege", "view") or r.darf("rechnungen", "view"):
        zeilen = []
        for re_ in Rechnung.objects.filter(verein=request.verein, mitglied=m, typ__in=["beitrag", "individuell"]) \
                .order_by("-datum"):
            zeilen.append({"url": reverse("rechnung_detail", args=[re_.pk]), "zellen": [
                re_.jahr or re_.datum.year, wert(re_, "betrag"), re_.nummer, wert(re_, "bezahlt_summe"),
                re_.get_status_display()]})
        abschnitte.append({"titel": "Beiträge & Zahlungen", "spalten": ["Jahr", "Beitrag", "Rechnung", "Bezahlt",
                                                                       "Status"], "zeilen": zeilen, "add_url": None})
    abschnitte.append(abschnitt(request, "Funktionen", m.funktionen.all(), ("funktion", "von", "bis"),
                                "mitgliedfunktion_add", {"mitglied": m.pk}))
    if r.darf("ehrungen", "view"):
        abschnitte.append(abschnitt(request, "Ehrungen", Ehrung.objects.filter(mitglied=m),
                                    ("art", "datum", "anlass"), "ehrung_add", {"mitglied": m.pk}))
    if r.darf("dokumente", "view"):
        abschnitte.append(abschnitt(request, "Dokumente", m.dokumente.all(),
                                    ("titel", "version", "kategorie", "erstellt"), "dokument_add", {"mitglied": m.pk}))
    if r.darf("verleih", "view"):
        abschnitte.append(abschnitt(request, "Verleih", Verleih.objects.filter(entleiher=m),
                                    ("gegenstand", "von", "bis", "status")))
    if r.darf("spenden", "view"):
        abschnitte.append(abschnitt(request, "Spenden", Spende.objects.filter(spender=m),
                                    ("datum", "art", "betrag")))
    if r.darf("aufwand", "view"):
        abschnitte.append(abschnitt(request, "Aufwandsentschädigungen",
                                    Aufwandsentschaedigung.objects.filter(empfaenger=m),
                                    ("datum", "art", "betrag", "status")))
    if r.darf("audit", "view"):
        abschnitte.append(abschnitt(request, "Änderungshistorie", AuditLog.objects.filter(
            verein=request.verein, modell="Mitglied", objekt_id=str(m.pk))[:20],
            ("zeit", "user_name", "aktion")))
    return {"aktionen": aktionen, "abschnitte": abschnitte}


def _dict(o, ausschluss=("verein", "openslides_initialpasswort")):
    d = {}
    for f in o._meta.concrete_fields:
        if f.name in ausschluss:
            continue
        v = getattr(o, f.attname)
        d[str(f.verbose_name)] = None if v in (None, "") else str(v)
    return d


@login_required
def mitglied_export(request, pk):
    """DSGVO-Auskunft: alle zu einem Mitglied gespeicherten Daten als JSON."""
    if request.verein is None or not request.rechte.darf("mitglieder", "view"):
        raise PermissionDenied
    from apps.allowances.models import Aufwandsentschaedigung
    from apps.donations.models import Spende
    from apps.finance.models import Rechnung, Zahlung
    from apps.honors.models import Ehrung

    m = get_object_or_404(Mitglied, pk=pk, verein=request.verein)
    daten = {
        "mitglied": _dict(m),
        "abteilungen": [str(a) for a in m.abteilungen.all()],
        "funktionen": [_dict(x) for x in m.funktionen.all()],
        "ehrungen": [_dict(x) for x in Ehrung.objects.filter(mitglied=m)],
        "rechnungen": [_dict(x) for x in Rechnung.objects.filter(mitglied=m)],
        "zahlungen": [_dict(x) for x in Zahlung.objects.filter(rechnung__mitglied=m)],
        "spenden": [_dict(x) for x in Spende.objects.filter(spender=m)],
        "aufwandsentschaedigungen": [_dict(x) for x in Aufwandsentschaedigung.objects.filter(empfaenger=m)],
        "dokumente": [_dict(x) for x in m.dokumente.all()],
    }
    r = HttpResponse(json.dumps(daten, ensure_ascii=False, indent=2), content_type="application/json; charset=utf-8")
    r["Content-Disposition"] = f'attachment; filename="mitglied-{m.mitgliedsnummer}-auskunft.json"'
    return r


@login_required
@require_POST
def mitglied_anonymisieren(request, pk):
    if request.verein is None or not request.rechte.darf("mitglieder", "delete"):
        raise PermissionDenied
    m = get_object_or_404(Mitglied, pk=pk, verein=request.verein)
    if m.openslides_user_id:
        from apps.openslides import services as os_services
        from apps.openslides.client import OpenSlidesFehler
        v = os_services.verbindung_oder_none(request.verein)
        if v is not None:
            try:
                os_services.mitglied_anonymisieren(v, m)
            except OpenSlidesFehler as e:
                messages.warning(request, f"OpenSlides-Konto konnte nicht angepasst werden ({e}) - bitte dort "
                                          "Name/E-Mail manuell prüfen und löschen.")
        else:
            messages.warning(request, "Mitglied hat ein OpenSlides-Konto, die Anbindung ist aber nicht "
                                      "eingerichtet - bitte Name/E-Mail dort manuell prüfen und löschen.")
        m.openslides_username, m.openslides_initialpasswort = "", ""
    m.vorname, m.nachname = "Anonymisiert", f"#{m.mitgliedsnummer}"
    for f in ("anrede", "strasse", "plz", "ort", "email", "telefon", "mobil", "kontoinhaber", "iban", "bic",
              "mandatsreferenz", "notizen"):
        setattr(m, f, "")
    m.geburtsdatum = m.mandatsdatum = None
    if m.foto:
        m.foto.delete(save=False)
    m.foto = ""
    if m.status in ("aktiv", "ruhend"):
        m.status, m.austrittsdatum = "ausgetreten", m.austrittsdatum or date.today()
    m.save()
    m.dokumente.all().delete()
    messages.success(request, "Mitglied anonymisiert.")
    return redirect("mitglied_detail", pk=m.pk)


# ---------------------------------------------------------------- Import / Export
def _log(request, aktion, text):
    from apps.core.models import AuditLog
    AuditLog.objects.create(verein=request.verein, user=request.user, user_name=request.user.get_username(),
                            ip=request.META.get("REMOTE_ADDR"), modell="Mitglied", objekt_id="0", objekt_repr=text[:200],
                            aktion=aktion, aenderungen={})


@login_required
def mitglieder_import(request):
    from django.shortcuts import render
    from . import importer
    if request.verein is None or not request.rechte.darf("mitglieder", "add"):
        raise PermissionDenied
    bericht = None
    if request.method == "POST":
        datei = request.FILES.get("datei")
        if not datei:
            messages.error(request, "Bitte eine Datei auswählen.")
        else:
            try:
                bericht = importer.importieren(
                    request.verein, datei.name, datei.read(), testlauf=bool(request.POST.get("testlauf")),
                    aktualisieren=bool(request.POST.get("aktualisieren")), neu_anlegen=bool(request.POST.get("neu_anlegen")))
                if not bericht["testlauf"]:
                    _log(request, "importiert", f"Import {datei.name}: {bericht['neu']} neu, {bericht['aktualisiert']} aktualisiert, "
                                                f"{len(bericht['fehler'])} Fehler")
            except ValueError as e:
                messages.error(request, str(e))
    return render(request, "members/import.html", {"titel": "Mitglieder importieren", "bericht": bericht,
                                                   "post": request.POST if request.method == "POST" else {}})


@login_required
def mitglieder_import_vorlage(request):
    if request.verein is None or not request.rechte.darf("mitglieder", "add"):
        raise PermissionDenied
    from io import BytesIO

    from openpyxl import Workbook

    from .tabellen import SPALTEN_ANZEIGE
    wb = Workbook()
    ws = wb.active
    ws.title = "Mitglieder"
    ws.append([label for _, label in SPALTEN_ANZEIGE])
    beispiel = {"vorname": "Erika", "nachname": "Beispiel", "anrede": "Frau", "geburtsdatum": "01.02.1980",
                "eintrittsdatum": "01.01.2015", "status": "aktiv", "mitgliedsart": "Aktiv", "strasse": "Hauptstr. 1",
                "plz": "35683", "ort": "Dillenburg", "email": "erika@example.org", "zahlungsart": "Überweisung",
                "abteilungen": "Vorstand, Jugend"}
    ws.append([beispiel.get(f, "") for f, _ in SPALTEN_ANZEIGE])
    for i, _ in enumerate(SPALTEN_ANZEIGE, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 18
    hilfe = wb.create_sheet("Hinweise")
    for z in ["Pflichtspalten: Vorname, Nachname.",
              "Leere Mitgliedsnummer = automatisch vergeben. Vorhandene Nummer = bestehendes Mitglied wird aktualisiert.",
              "Ohne Mitgliedsnummer wird über Vorname + Nachname + Geburtsdatum abgeglichen.",
              "Leere Zellen überschreiben keine vorhandenen Daten.",
              "Datum: TT.MM.JJJJ. Status: aktiv, ruhend, ausgetreten, verstorben. Anrede: Herr, Frau, Divers, Firma.",
              "Mitgliedsart/Abteilungen müssen existieren (oder Option „unbekannte anlegen“ beim Import).",
              "Zeile 2 ist ein Beispiel und sollte gelöscht werden."]:
        hilfe.append([z])
    buf = BytesIO()
    wb.save(buf)
    r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = 'attachment; filename="mitglieder-import-vorlage.xlsx"'
    return r


@login_required
def mitglieder_export(request):
    """Vollständiger Export aller Mitgliederdaten als Excel/CSV (Bankdaten nur mit Beitragsrecht)."""
    import csv
    from io import BytesIO, StringIO

    from apps.core.crud import _sicher

    from .tabellen import SPALTEN_ANZEIGE
    if request.verein is None or not request.rechte.darf("mitglieder", "view"):
        raise PermissionDenied
    mit_bank = bool(request.GET.get("bank")) and request.rechte.darf("beitraege", "view")
    bank_felder = {"zahlungsart", "kontoinhaber", "iban", "bic", "mandatsreferenz", "mandatsdatum", "individueller_beitrag"}
    spalten = [(f, l) for f, l in SPALTEN_ANZEIGE if mit_bank or f not in bank_felder]
    qs = Mitglied.objects.filter(verein=request.verein).select_related("mitgliedsart", "familie").prefetch_related("abteilungen")
    if request.GET.get("status") in ("aktiv", "ruhend", "ausgetreten", "verstorben"):
        qs = qs.filter(status=request.GET["status"])
    qs = qs.order_by("mitgliedsnummer")

    def zelle(m, f):
        if f == "abteilungen":
            return ", ".join(a.name for a in m.abteilungen.all())
        if f == "status":
            return m.status
        if f == "anrede":
            return m.get_anrede_display() if m.anrede else ""
        if f == "zahlungsart":
            return m.get_zahlungsart_display()
        if f == "ist_familienzahler":
            return "ja" if m.ist_familienzahler else ""
        v = getattr(m, f)
        if hasattr(v, "strftime"):
            return v.strftime("%d.%m.%Y")
        return "" if v is None else str(v).replace(".", ",") if f == "individueller_beitrag" else ("" if v is None else str(v))
    zeilen = [[_sicher(zelle(m, f)) for f, _ in spalten] for m in qs]
    kopf = [l for _, l in spalten]
    _log(request, "exportiert", f"Mitgliederexport ({len(zeilen)} Datensätze{', mit Bankdaten' if mit_bank else ''})")
    name = f"mitglieder-{date.today():%Y%m%d}"
    if request.GET.get("format") == "csv":
        out = StringIO()
        w = csv.writer(out, delimiter=";")
        w.writerow(kopf)
        w.writerows(zeilen)
        r = HttpResponse("\ufeff" + out.getvalue(), content_type="text/csv; charset=utf-8")
        r["Content-Disposition"] = f'attachment; filename="{name}.csv"'
        return r
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Mitglieder"
    ws.append(kopf)
    for z in zeilen:
        ws.append(z)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return r
